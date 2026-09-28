"""The opening price list (Sales step S1): the paper list typed into a system template (Tires / Batteries / Mags tabs, the paper list's
columns: brand, size or model, qty = stock on hand, remarks, cost), uploaded, checked row by row, then confirmed by the owner.
Nothing is saved from an upload that has a problem."""
import io
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from inventory.models import Category, Item
from inventory.stock import item_key

SHEETS = {"Tires": Category.TIRE, "Batteries": Category.BATTERY, "Mags": Category.MAGS}
COLUMNS = {Category.TIRE: ["Brand", "Size", "Qty", "Remarks", "Cost"],
           Category.BATTERY: ["Brand", "Model", "Qty", "Remarks", "Cost"],
           Category.MAGS: ["Brand", "Model", "Qty", "Remarks", "Cost"]}
HEADER_WORDS = {"brand": "brand", "size": "size_model", "model": "size_model", "size / model": "size_model", "type": "size_model",
                "qty": "qty", "quantity": "qty", "pcs": "qty", "on hand": "qty", "remarks": "remark", "remark": "remark",
                "notes": "remark", "cost": "cost", "unit cost": "cost"}
MAX_ROWS = 5000
NAVY = "042C53"


def build_template() -> bytes:
    wb = Workbook()
    how = wb.active
    how.title = "How to fill"
    for i, line in enumerate([
        "Slicks ERP -- opening price list",
        "",
        "Copy the paper price list into the three tabs: Tires, Batteries, Mags (one row per item, same columns as the paper list).",
        "Brand and Size (or Model) are required. Qty = pieces on hand today (0 if none). Cost = cost per piece, e.g. 3150 or 3,150.00.",
        "Remarks: only when needed, e.g. 'old stock 2023' -- also use it when the same brand and size is on two rows at different costs.",
        "Leave the headings as they are. Blank rows are skipped. Upload the file on the ERP: Items and stock -> Opening list.",
        "The owner checks the totals and confirms before it becomes the opening stock.",
    ], start=1):
        how.cell(row=i, column=1, value=line)
    how["A1"].font = Font(bold=True, size=14, color=NAVY)
    how.column_dimensions["A"].width = 120
    for title, cat in SHEETS.items():
        ws = wb.create_sheet(title)
        for col, head in enumerate(COLUMNS[cat], start=1):
            c = ws.cell(row=1, column=col, value=head)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(horizontal="center")
        for col, width in zip("ABCDE", (22, 22, 8, 30, 14)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        qty = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                             error="Whole pieces, 0 or more.", showErrorMessage=True)
        cost = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                              error="Cost per piece, 0 or more.", showErrorMessage=True)
        ws.add_data_validation(qty)
        ws.add_data_validation(cost)
        qty.add(f"C2:C{MAX_ROWS}")
        cost.add(f"E2:E{MAX_ROWS}")
        for row in range(2, MAX_ROWS + 1):
            ws.cell(row=row, column=5).number_format = "#,##0.00"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _text(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def parse(data: bytes) -> dict:
    """Read an uploaded template. Returns {'rows', 'errors', 'items_count', 'pieces', 'value'}; rows are only usable with no errors."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        return {"rows": [], "errors": ["This isn't an Excel file the system can read. Use the template from Items and stock."],
                "items_count": 0, "pieces": 0, "value": Decimal("0")}
    rows, errors, keys = [], [], {}
    found = False
    for ws in wb.worksheets:
        cat = next((c for t, c in SHEETS.items() if t.lower() == ws.title.strip().lower()), None)
        if cat is None:
            continue
        found = True
        it = ws.iter_rows(values_only=True)
        header = next(it, None) or ()
        cols = {HEADER_WORDS.get(_text(h).lower()): i for i, h in enumerate(header) if HEADER_WORDS.get(_text(h).lower())}
        missing = [n for n in ("brand", "size_model", "qty", "cost") if n not in cols]
        if missing:
            errors.append(f"{ws.title}: the headings changed -- missing {', '.join(missing).replace('size_model', 'size / model')}.")
            continue
        for n, values in enumerate(it, start=2):
            get = lambda name: values[cols[name]] if name in cols and cols[name] < len(values) else None
            brand, size, remark = _text(get("brand")), _text(get("size_model")), _text(get("remark"))
            qty, cost = get("qty"), get("cost")
            if not brand and not size and qty in (None, "") and cost in (None, ""):
                continue
            where = f"{ws.title} row {n}"
            if not brand or not size:
                errors.append(f"{where}: brand and {'size' if cat == Category.TIRE else 'model'} are both needed.")
                continue
            try:
                q = Decimal(str(qty if qty not in (None, "") else 0).replace(",", ""))
                c = Decimal(str(cost).replace(",", "")) if cost not in (None, "") else None
            except InvalidOperation:
                errors.append(f"{where}: qty and cost must be numbers.")
                continue
            if c is None:
                errors.append(f"{where}: the cost is missing.")
                continue
            if q < 0 or q != q.to_integral_value():
                errors.append(f"{where}: qty must be whole pieces, 0 or more.")
                continue
            if c < 0:
                errors.append(f"{where}: the cost can't be negative.")
                continue
            key = item_key(cat, brand, size, remark)
            if key in keys:
                errors.append(f"{where}: the same item as {keys[key]} -- merge them, or add a remark to tell them apart.")
                continue
            keys[key] = where
            existing = Item.objects.filter(key=key).first()
            if existing and existing.moves.exists():
                errors.append(f"{where}: {existing.code} {existing.description} already has stock movements -- leave it out.")
                continue
            rows.append({"category": cat, "brand": brand, "size_model": size, "remark": remark, "qty": int(q),
                         "cost": str(c.quantize(Decimal("0.01")))})
    if not found:
        errors.append("No Tires, Batteries or Mags tab found. Use the template from Items and stock.")
    elif not rows and not errors:
        errors.append("The file has no items.")
    pieces = sum(r["qty"] for r in rows)
    value = sum((r["qty"] * Decimal(r["cost"]) for r in rows), Decimal("0")).quantize(Decimal("0.01"))
    return {"rows": rows, "errors": errors, "items_count": len(rows), "pieces": pieces, "value": value}
