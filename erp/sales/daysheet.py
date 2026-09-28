"""The Excel day sheet (S2). The owner approved a macro pop-up form; building a macro workbook needs an Excel security setting changed
(the owner's decision), so this first version has NO macros: dropdowns (Microsoft 365 narrows them as you type) and formulas give the
same result -- invoice numbers from the booklet in use, customers, items with cost; cost, amount and profit fill in; the day's totals
add up. Downloaded fresh each day with today's lists; uploaded at closing; every line checked; saved all together or not at all.

Sheet "Day": B1 = date; row 2 = the day's totals; row 4 = headings; rows 5.. = one row per item line in the notebook's order
(qty, item, price, cost, amount, profit, credit customer) plus invoice no., payment and the payment details.
Sheet "Cancelled": spoiled / cancelled invoice numbers with the reason. Hidden sheets hold the lists."""
import io
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core.models import Party
from engines.clock import manila_today
from inventory.models import Item
from sales import rules
from sales.models import PaymentType

ROWS = 150                       # entry rows on the day sheet
FIRST = 5                        # first entry row
HEAD = ["Qty", "Item", "Price each", "Cost each", "Amount", "Profit", "Customer", "Invoice no.", "Payment",
        "Check bank", "Check no.", "Check date", "Reference / approval", "Below-cost reason"]
NAVY = "042C53"
PAY_LABEL = {p.value: p.label for p in PaymentType}
PAY_BY_LABEL = {p.label.lower(): p.value for p in PaymentType} | {p.value: p.value for p in PaymentType} | {"gcash": "ewallet", "maya": "ewallet", "credit": "charge"}


def item_choice(i: Item) -> str:
    return f"{i.code} · {i.description}"


def _customers():
    return list(Party.objects.filter(is_active=True).filter(is_customer=True).exclude(is_walk_in=True).order_by("name")) \
        + list(Party.objects.filter(is_active=True, is_affiliate=True, is_customer=False).order_by("name"))


def build_template(day: date | None = None) -> bytes:
    day = day or manila_today()
    b = rules.current_booklet()
    numbers = [b.fmt(n) for n in rules.unused_numbers(b)] if b else []
    items = list(Item.objects.filter(is_active=True).order_by("category", "brand", "size_model", "remark"))
    customers = _customers()

    wb = Workbook()
    ws = wb.active
    ws.title = "Day"
    ws["A1"], ws["B1"] = "Date", day
    ws["B1"].number_format = "mmm d, yyyy"
    ws["D1"] = f"Booklet in use: {b}" if b else "No booklet in use -- ask the owner to register one"
    ws["A2"] = "Day total"
    ws["A2"].font = Font(bold=True)
    last = FIRST + ROWS - 1
    ws["B2"] = f'=SUM(A{FIRST}:A{last})&" pcs"'
    ws["E2"] = f"=SUM(E{FIRST}:E{last})"
    ws["F2"] = f"=SUM(F{FIRST}:F{last})"
    for c in ("E2", "F2"):
        ws[c].number_format = "#,##0.00"
        ws[c].font = Font(bold=True)
    for col, head in enumerate(HEAD, start=1):
        cell = ws.cell(row=4, column=col, value=head)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for col, width in zip("ABCDEFGHIJKLMN", (6, 38, 12, 12, 12, 12, 26, 12, 16, 12, 12, 12, 18, 24)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"
    grey = PatternFill("solid", fgColor="F1EFE8")
    for r in range(FIRST, last + 1):
        # INDEX / MATCH (not XLOOKUP): understood by every Excel version and Google Sheets; checked in real Excel 2026-09-29
        ws[f"D{r}"] = f'=IFERROR(INDEX(Lists!$B:$B,MATCH(B{r},Lists!$A:$A,0)),"")'
        ws[f"E{r}"] = f'=IF(AND(A{r}<>"",C{r}<>""),A{r}*C{r},"")'
        ws[f"F{r}"] = f'=IF(AND(E{r}<>"",D{r}<>""),E{r}-A{r}*D{r},"")'
        for c in "CDEF":
            ws[f"{c}{r}"].number_format = "#,##0.00"
        ws[f"L{r}"].number_format = "mmm d, yyyy"
        for c in "DEF":
            ws[f"{c}{r}"].fill = grey

    lists = wb.create_sheet("Lists")
    lists.append(["Item", "Avg cost", "On hand"])
    for i in items:
        lists.append([item_choice(i), float(i.avg_cost), i.on_hand])
    for r, c in enumerate(customers, start=2):
        lists.cell(row=r, column=5, value=c.name)
        lists.cell(row=r, column=6, value="credit" if c.credit_allowed else "cash only")
    lists["E1"], lists["F1"] = "Customer", "Credit"
    for r, n in enumerate(numbers, start=2):
        lists.cell(row=r, column=8, value=n)
    lists["H1"] = "Invoice no."
    for r, p in enumerate(PaymentType, start=2):
        lists.cell(row=r, column=10, value=p.label)
    lists["J1"] = "Payment"
    lists.sheet_state = "hidden"

    def dv(formula, cells, msg):
        v = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True, error=msg)
        ws.add_data_validation(v)
        v.add(cells)
        return v
    dv(f"=Lists!$A$2:$A${max(len(items), 1) + 1}", f"B{FIRST}:B{last}", "Pick the item from the list (ask the owner to add a new one).")
    dv(f"=Lists!$E$2:$E${max(len(customers), 1) + 1}", f"G{FIRST}:G{last}", "Pick a customer from the list, or leave it empty for a walk-in.")
    dv(f"=Lists!$H$2:$H${max(len(numbers), 1) + 1}", f"H{FIRST}:H{last}", "Pick an unused number of the booklet in use.")
    dv(f"=Lists!$J$2:$J${len(PaymentType) + 1}", f"I{FIRST}:I{last}", "Pick how it was paid.")
    q = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True, showErrorMessage=True,
                       error="Whole pieces, at least 1.")
    ws.add_data_validation(q)
    q.add(f"A{FIRST}:A{last}")

    can = wb.create_sheet("Cancelled")
    can.append(["Invoice no.", "Reason (spoiled, wrong amount written, ...)"])
    for cell in can[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    can.column_dimensions["A"].width, can.column_dimensions["B"].width = 14, 60
    v = DataValidation(type="list", formula1=f"=Lists!$H$2:$H${max(len(numbers), 1) + 1}", allow_blank=True)
    can.add_data_validation(v)
    v.add("A2:A40")

    how = wb.create_sheet("How to fill", 0)
    for r, line in enumerate([
        "Slicks ERP -- day sheet",
        "Day tab: one row per item sold, like the notebook (qty, item, price). Cost, amount, profit and the day's total fill in by themselves.",
        "Invoice no.: pick it from the list (unused numbers of the booklet in use). Several items on one invoice = the same number on several rows.",
        "Customer: leave empty for a walk-in (cash). Charge (credit) only for customers the owner allowed credit.",
        "Payment: a check needs the bank, check no. and check date; GCash / Maya, transfer and card need the reference / approval code.",
        "Below cost: allowed -- write a short reason (e.g. 'keep customer'); the owner reviews it.",
        "Cancelled tab: a spoiled or cancelled paper invoice -- its number and the reason.",
        "At closing: Day sheet page on the ERP -> upload. Nothing is saved until every line is fine; then Save the day.",
    ], start=1):
        how.cell(row=r, column=1, value=line)
    how["A1"].font = Font(bold=True, size=14, color=NAVY)
    how.column_dimensions["A"].width = 130
    wb.active = wb.sheetnames.index("Day")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---- reading an upload --------------------------------------------------------------------------------------------------------------

def _t(v) -> str:
    return " ".join(str(v).split()) if v not in (None, "") else ""


def _day(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _t(v)
    for fmt in ("%Y-%m-%d", "%b %d, %Y", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parse(data: bytes, file_name: str = "") -> dict:
    """Read and check an uploaded day sheet. Returns {'date', 'sales': [serialisable], 'cancel': [(no, reason)], 'errors', 'warnings',
    'total', 'profit', 'count'}. Nothing is saved here."""
    out = {"date": None, "sales": [], "cancel": [], "errors": [], "warnings": [], "total": Decimal("0"), "profit": Decimal("0"), "count": 0}
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        out["errors"].append("This isn't an Excel file the system can read. Use the day sheet downloaded from the ERP.")
        return out
    if "Day" not in wb.sheetnames:
        out["errors"].append("No Day tab found. Use the day sheet downloaded from the ERP.")
        return out
    ws = wb["Day"]
    day = _day(ws["B1"].value)
    if day is None:
        out["errors"].append("Day tab: the date in B1 is missing or unreadable.")
        return out
    out["date"] = day
    items = {i.code: i for i in Item.objects.all()}
    customers = {c.name.lower(): c for c in Party.objects.all()}
    walk_in = Party.objects.get(is_walk_in=True)
    groups: "OrderedDict[str, dict]" = OrderedDict()
    for r in range(FIRST, ws.max_row + 1):
        v = lambda col: ws[f"{col}{r}"].value
        qty, item_text, price, number = v("A"), _t(v("B")), v("C"), _t(v("H"))
        if qty in (None, "") and not item_text and price in (None, "") and not number:
            continue
        where = f"Row {r}"
        if not number:
            out["errors"].append(f"{where}: pick the invoice number.")
            continue
        code = item_text.split("·")[0].strip().upper()
        item = items.get(code)
        if item is None:
            out["errors"].append(f"{where}: pick the item from the list ({item_text or 'empty'}).")
            continue
        cust_text = _t(v("G"))
        customer = customers.get(cust_text.lower()) if cust_text else walk_in
        if customer is None:
            out["errors"].append(f"{where}: {cust_text} is not on the customer list.")
            continue
        pay = PAY_BY_LABEL.get(_t(v("I")).lower())
        if pay is None:
            out["errors"].append(f"{where}: pick how invoice {number} was paid.")
            continue
        head = {"customer": customer.pk, "payment": pay, "check_bank": _t(v("J")), "check_number": _t(v("K")),
                "check_date": (_day(v("L")).isoformat() if _day(v("L")) else ""), "payment_ref": _t(v("M"))}
        g = groups.get(number)
        if g is None:
            groups[number] = g = {"number": number, "rows": [], "head": head, "lines": []}
        else:
            for k, val in head.items():
                if val and g["head"][k] and val != g["head"][k]:
                    out["errors"].append(f"{where}: invoice {number} has a different {k.replace('_', ' ')} than on row {g['rows'][0]}.")
                elif val and not g["head"][k]:
                    g["head"][k] = val
        g["rows"].append(r)
        g["lines"].append({"item": item.pk, "qty": str(qty if qty is not None else ""), "price": str(price if price is not None else ""),
                           "reason": _t(v("N"))})
    if "Cancelled" in wb.sheetnames:
        for r, row in enumerate(wb["Cancelled"].iter_rows(min_row=2, values_only=True), start=2):
            no, reason = (_t(row[0]) if row else ""), (_t(row[1]) if row and len(row) > 1 else "")
            if not no and not reason:
                continue
            if not no or not reason:
                out["errors"].append(f"Cancelled tab row {r}: give both the invoice number and the reason.")
                continue
            if no in groups:
                out["errors"].append(f"Cancelled tab row {r}: {no} is also on the Day tab.")
                continue
            out["cancel"].append((no, reason))
    if not groups and not out["cancel"]:
        out["errors"].append("The Day tab has no sales.")
    for g in groups.values():
        s = to_input(g, day)
        try:
            chk = rules.check_sale(s)
        except rules.SaleError as e:
            out["errors"].append(f"Row {', '.join(map(str, g['rows']))}: {e}")
            continue
        out["warnings"] += chk["warnings"]
        out["total"] += chk["total"]
        out["profit"] += chk["profit"]
        out["count"] += 1
        out["sales"].append(g)
    for no, reason in out["cancel"]:
        try:
            n = rules._number(no)
            b = rules._booklet_in_use_for(n)
            if rules.Sale.objects.filter(number=n).exists() or rules.CancelledNumber.objects.filter(number=n).exists():
                out["errors"].append(f"Cancelled tab: {no} is already used or cancelled.")
        except rules.SaleError as e:
            out["errors"].append(f"Cancelled tab: {e}")
    b = rules.current_booklet()
    if b and not out["errors"]:
        mine = {rules._number(g["number"]) for g in out["sales"]} | {rules._number(no) for no, _ in out["cancel"]}
        used = rules.used_numbers(b) | mine
        taken = used | rules.cancelled_numbers(b) | mine
        skipped = [b.fmt(n) for n in range(b.first_number, max(used) if used else b.first_number) if n not in taken]
        if skipped:
            out["warnings"].append(f"Skipped number(s) {', '.join(skipped)} -- neither used nor cancelled; add them to the Cancelled tab "
                                   "if they were spoiled (the owner sees them on the gap list).")
    return out


def to_input(g: dict, day: date) -> rules.SaleInput:
    h = g["head"]
    lines = [rules.SaleLineInput(Item.objects.filter(pk=x["item"]).first(), x["qty"], x["price"], x.get("reason", "")) for x in g["lines"]]
    return rules.SaleInput(number=g["number"], sale_date=day, customer=Party.objects.filter(pk=h["customer"]).first(),
                           payment_type=h["payment"], lines=lines, check_bank=h["check_bank"], check_number=h["check_number"],
                           check_date=date.fromisoformat(h["check_date"]) if h["check_date"] else None, payment_ref=h["payment_ref"])


def save_day(parsed: dict, file_name: str) -> "rules.DayUpload":
    """All the day's sales and cancelled numbers in ONE locked save -- all or nothing (re-checked inside the lock)."""
    from django.db import transaction

    from core import audit
    from engines.acting_user import acting_user

    day = date.fromisoformat(parsed["date"]) if isinstance(parsed["date"], str) else parsed["date"]
    with transaction.atomic():
        up = rules.DayUpload.objects.create(file_name=file_name[:200], sale_date=day, sales_count=len(parsed["sales"]),
                                            total=Decimal(str(parsed["total"])), profit=Decimal(str(parsed["profit"])),
                                            uploaded_by=acting_user())
        for no, reason in parsed["cancel"]:
            rules.cancel_number(no, reason)
        for g in parsed["sales"]:
            rules.record_sale(to_input(g, day), source=rules.SaleSource.UPLOAD, upload=up)
        audit.log("SALES_DAY_UPLOADED", "sales_dayupload", up.pk,
                  new=f"{up.file_name}: {up.sale_date} {up.sales_count} sale(s) PHP {up.total:,.2f}, profit PHP {up.profit:,.2f}, "
                      f"{len(parsed['cancel'])} number(s) cancelled")
    return up
