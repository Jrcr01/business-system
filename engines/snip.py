"""Snip from a document (owner, 2026-10-16: like DataSnipper, inside the ERP; mock-ups docs/mockups/aa_snip.html and ab_snip_demo.html). A PDF or
photo is kept beside the database; the page is shown as a picture; the person drags a box; the words inside the box come out as a table or a text.

How a box is read, in this order (the accuracy comes from here, not from the OCR engine):
  1. a PDF with real text: the characters and their positions on the page, exact (pdfplumber);
     a. a ruled table inside the box (lines drawn on the page): pdfplumber's table reader on the crop;
     b. no lines: words that line up at the same x make a column, words at the same y make a row (the position method, as the demo);
  2. a scan or a photo (or a PDF page with no text): the box is cut out, enlarged, OCR'd (Tesseract) into words with positions, then 1b. A cell the OCR
     cannot read, or is unsure of, is left for the person to type (owner: "I am the one should enter it").
The document is kept only while it is being snipped (a list of many pages needs it between pages): it is deleted the moment its rows are imported,
or after a day if never used (owner, 2026-10-16: "no need to store the document"). Nothing leaves the server."""
from __future__ import annotations

import io
import json
import re
import secrets
import time
from pathlib import Path

from engines.clock import manila_now
from engines.common import get_conn, write_tx

KINDS = ("price", "loan", "soa", "card")                # card: a credit card statement (2026-10-17), owner_id None
EXTS = {".pdf": "pdf", ".jpg": "jpg", ".jpeg": "jpg", ".png": "png", ".webp": "png"}
MAX_BYTES = 25 * 1024 * 1024
RENDER_SCALE = 2.0                 # the page picture: 2x the PDF's points (sharp on a laptop, readable on a phone)
OCR_SCALE = 3                      # the crop is enlarged this much before OCR
OCR_DOUBT = 70                     # Tesseract's confidence under this: the cell is "doubtful" (shown orange, to confirm)
STAMP = "%Y-%m-%d %H:%M:%S"


def folder() -> Path:
    """The documents while they are being snipped, beside the database, never in git; deleted after use (owner, 2026-10-16)."""
    import engines.db_engine as d
    p = Path(d.DB_PATH).resolve().parent / "snip_files"
    p.mkdir(parents=True, exist_ok=True)
    return p


# --- the files ----------------------------------------------------------------------------------------------------------------------------------------

def keep(kind: str, owner_id: int | None, filename: str, data: bytes) -> tuple[bool, str, str | None]:
    """Stores an uploaded document. Returns (ok, message, token)."""
    if kind not in KINDS:
        return False, "Unknown kind of document.", None
    ext = EXTS.get(Path(filename or "").suffix.lower())
    if not ext:
        return False, "Choose a PDF or a picture (JPG, PNG).", None
    if not data:
        return False, "The file is empty.", None
    if len(data) > MAX_BYTES:
        return False, "The file is larger than 25 MB. Is it the right file?", None
    sweep()
    token = secrets.token_hex(12)
    path = folder() / f"{token}.{ext}"
    path.write_bytes(data)
    pages, has_text = _inspect(path, ext)
    if pages == 0:
        path.unlink(missing_ok=True)
        return False, "The file could not be opened as a PDF or a picture.", None
    with write_tx() as conn:
        conn.execute("""INSERT INTO snip_files (token, kind, owner_id, filename, ext, pages, has_text, uploaded_by, created_at)
                        VALUES (?, ?, ?, ?, ?, ?, ?, acting_user(), ?)""",
                     (token, kind, owner_id, (filename or "")[:200], ext, pages, int(has_text), manila_now().strftime(STAMP)))
    return True, f"{filename} kept: {pages} page{'s' if pages != 1 else ''}" + ("" if has_text else " (a picture: read by OCR)") + ".", token


def _inspect(path: Path, ext: str) -> tuple[int, bool]:
    try:
        if ext == "pdf":
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                n = len(pdf.pages)
                text = any(len((p.extract_text() or "").strip()) > 20 for p in pdf.pages[:3])
            return n, text
        from PIL import Image
        with Image.open(path) as im:
            im.verify()
        return 1, False
    except Exception:                                   # noqa: BLE001
        return 0, False


def discard(token: str) -> None:
    """The document, its page pictures and its snips are deleted: after the rows were imported, or when it was never used."""
    info = file_info(token)
    if not info:
        return
    for p in folder().glob(f"{token}*"):
        try:
            p.unlink()
        except OSError:
            pass
    with write_tx() as conn:
        conn.execute("DELETE FROM snips WHERE file_id = ?", (info["file_id"],))
        conn.execute("DELETE FROM snip_files WHERE file_id = ?", (info["file_id"],))


def sweep(hours: int = 24) -> int:
    """Documents uploaded more than `hours` ago and never imported are deleted (run when a new one is opened)."""
    cutoff = (manila_now() - __import__("datetime").timedelta(hours=hours)).strftime(STAMP)
    with get_conn() as conn:
        old = [r[0] for r in conn.execute("SELECT token FROM snip_files WHERE created_at < ?", (cutoff,)).fetchall()]
    for token in old:
        discard(token)
    return len(old)


def file_info(token: str) -> dict | None:
    if not token or not re.fullmatch(r"[0-9a-f]{24}", token):
        return None
    with get_conn() as conn:
        row = conn.execute("SELECT file_id, token, kind, owner_id, filename, ext, pages, has_text, uploaded_by, created_at FROM snip_files WHERE token = ?", (token,)).fetchone()
    if not row:
        return None
    keys = ("file_id", "token", "kind", "owner_id", "filename", "ext", "pages", "has_text", "uploaded_by", "created_at")
    out = dict(zip(keys, row))
    out["path"] = folder() / f"{out['token']}.{out['ext']}"
    return out


def files_for(kind: str, owner_id: int | None, limit: int = 12) -> list[dict]:
    with get_conn() as conn:
        rows = conn.execute("""SELECT token, filename, pages, has_text, uploaded_by, created_at, (SELECT COUNT(*) FROM snips s WHERE s.file_id = f.file_id)
                               FROM snip_files f WHERE kind = ? AND (owner_id = ? OR ? IS NULL) ORDER BY file_id DESC LIMIT ?""", (kind, owner_id, owner_id, limit)).fetchall()
    return [{"token": a, "filename": b, "pages": c, "has_text": bool(d), "by": e, "at": f, "snips": g} for a, b, c, d, e, f, g in rows]


# --- the page picture and its words -------------------------------------------------------------------------------------------------------------------

# --- straightening (owner, 2026-10-17: mock-up ap_snip_rows_columns, answer 1 A) ----------------------------------------------------------------------
# A photo taken at a slant has rows that slope, so a straight cut takes the bottom of one row and the top of the next. The page is turned until the
# table's own lines run level, once, before anything is shown or read; the person can nudge it (set_angle). A PDF with real text is never turned.

def straightens(info: dict) -> bool:
    return not (info["ext"] == "pdf" and info["has_text"])


def _angle_path(token: str, page: int) -> Path:
    return folder() / f"{token}.p{page}.angle"


def _detect_angle(img) -> float:
    """The turn (degrees, counter-clockwise) that makes the page's lines of text and rules level: the angle whose row profile is sharpest."""
    import numpy as np
    from PIL import Image, ImageOps
    g = ImageOps.grayscale(img)
    if g.width > 900:
        g = g.resize((900, max(1, int(g.height * 900 / g.width))), Image.BILINEAR)
    g = ImageOps.autocontrast(g)

    def score(a: float) -> float:
        r = g.rotate(a, resample=Image.BILINEAR, fillcolor=255)
        dark = (np.asarray(r) < 140).sum(axis=1).astype(float)
        return float(np.sum(np.diff(dark) ** 2))
    best = max((a / 4 for a in range(-32, 33)), key=score)            # -8 .. +8 degrees, by a quarter
    best = max((best + a / 20 for a in range(-5, 6)), key=score)       # then by a twentieth
    return round(best, 2) if abs(best) >= 0.15 else 0.0


def page_angle(token: str, page: int) -> float:
    """The page's turn: found the first time the page is shown, then kept (and changed only by the person)."""
    info = file_info(token)
    if not info or not straightens(info):
        return 0.0
    p = _angle_path(token, page)
    if p.exists():
        try:
            return float(p.read_text(encoding="utf-8").strip() or 0)
        except ValueError:
            pass
    try:
        a = _detect_angle(_raw_picture(info, page, RENDER_SCALE))
    except Exception:                                   # noqa: BLE001 -- a picture that cannot be measured is shown as it is
        a = 0.0
    p.write_text(str(a), encoding="utf-8")
    return a


def set_angle(token: str, page: int, angle: float) -> float:
    """The person's own turn (−15° .. +15°): kept, and the page picture made again."""
    info = file_info(token)
    if not info or not straightens(info) or not 1 <= page <= info["pages"]:
        return 0.0
    a = round(max(-15.0, min(15.0, float(angle))), 2)
    _angle_path(token, page).write_text(str(a), encoding="utf-8")
    for old in folder().glob(f"{token}.p{page}.*jpg"):          # the pictures at the old turn (one may still be open while it is sent: left for discard)
        try:
            old.unlink()
        except OSError:
            pass
    return a


def _raw_picture(info: dict, page: int, scale: float):
    """The page as an RGB picture, not turned: a PDF page rendered at `scale`, a photo as taken (its EXIF turn applied)."""
    from PIL import Image, ImageOps
    if info["ext"] == "pdf":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(info["path"]))
        img = pdf[page - 1].render(scale=scale).to_pil().convert("RGB")
        pdf.close()                                              # else Windows keeps the file locked and it cannot be deleted after use
        return img
    return ImageOps.exif_transpose(Image.open(info["path"])).convert("RGB")


def _picture(info: dict, page: int, scale: float):
    """The page as the person sees it: turned by the page's angle (same size, white corners), so a box drawn on the screen is the box read."""
    from PIL import Image
    img = _raw_picture(info, page, scale)
    a = page_angle(info["token"], page) if straightens(info) else 0.0
    return img.rotate(a, resample=Image.BICUBIC, fillcolor=(255, 255, 255)) if a else img


def page_image(token: str, page: int) -> Path | None:
    """The page as a JPEG, made once and kept beside the file."""
    info = file_info(token)
    if not info or not 1 <= page <= info["pages"]:
        return None
    a = page_angle(token, page)
    out = folder() / (f"{token}.p{page}.jpg" if not a else f"{token}.p{page}.a{a:+.2f}.jpg")     # one picture per turn: a nudge never waits on a file in use
    if out.exists():
        return out
    try:
        img = _picture(info, page, RENDER_SCALE)
        img.save(out, "JPEG", quality=82)
    except Exception:                                   # noqa: BLE001
        return None
    return out


def page_size(token: str, page: int) -> tuple[float, float]:
    """The coordinate space of a page: PDF points for a PDF, pixels for a picture (what the boxes are measured in)."""
    info = file_info(token)
    if info["ext"] == "pdf":
        import pdfplumber
        with pdfplumber.open(str(info["path"])) as pdf:
            p = pdf.pages[page - 1]
            return float(p.width), float(p.height)
    from PIL import Image, ImageOps
    with Image.open(info["path"]) as im:
        im = ImageOps.exif_transpose(im)
        return float(im.width), float(im.height)


def _pdf_words(info: dict, page: int) -> list[dict]:
    import pdfplumber
    with pdfplumber.open(str(info["path"])) as pdf:
        p = pdf.pages[page - 1]
        return [{"t": w["text"], "x0": float(w["x0"]), "x1": float(w["x1"]), "y0": float(w["top"]), "y1": float(w["bottom"]), "conf": 100}
                for w in p.extract_words(keep_blank_chars=False, use_text_flow=False)]


# --- RapidOCR (owner, 2026-10-17: "setup AI" -> RapidOCR for snips) ------------------------------------------------------------------------------------
# On the owner's Ultra Plus battery photo it read 96 of 96 cells where Tesseract read 11 of 39 (scratch comparison, 2026-10-17). It runs on our own
# server (PaddleOCR's models via onnxruntime, ~100 MB once loaded, per worker that reads a picture); nothing leaves the server. It finds the text in
# the whole box itself -- its strength -- and each piece goes to the cell its centre falls in. Tesseract stays as the fallback when it cannot load.

RAPID_DOUBT = 85                   # RapidOCR's score (0-100) under this: the cell is shown orange, to confirm
_RAPID = {"engine": None, "failed": False}
import threading as _threading
_RAPID_LOCK = _threading.Lock()


def _rapid():
    """The RapidOCR engine, loaded once per process on first use; None when it cannot load (then Tesseract reads)."""
    if _RAPID["engine"] is None and not _RAPID["failed"]:
        with _RAPID_LOCK:
            if _RAPID["engine"] is None and not _RAPID["failed"]:
                try:
                    from rapidocr_onnxruntime import RapidOCR
                    _RAPID["engine"] = RapidOCR()
                except Exception:                       # noqa: BLE001 -- not installed, or its models missing
                    _RAPID["failed"] = True
    return _RAPID["engine"]


def _rapid_words(info: dict, page: int, box) -> list[dict] | None:
    """The text pieces RapidOCR finds inside the box, with their positions in the page's own coordinates; None when RapidOCR is not there."""
    eng = _rapid()
    if eng is None:
        return None
    import numpy as np
    from PIL import Image
    scale = 3.0 if info["ext"] == "pdf" else 1.0                 # a scanned PDF page rendered sharp; a photo as it is
    img = _picture(info, page, scale)
    x0, y0, x1, y1 = [v * scale for v in box]
    crop = img.crop((int(x0), int(y0), int(x1) + 1, int(y1) + 1))
    up = 2 if crop.width < 2400 else 1                            # small text is enlarged (the tested setting)
    if up != 1:
        crop = crop.resize((crop.width * up, crop.height * up), Image.LANCZOS)
    with _RAPID_LOCK:                                             # one read at a time per process: the server has one processor
        res, _ = eng(np.asarray(crop))
    out = []
    for pts, text, score in res or []:
        t = _clean_cell(text)
        if not t:
            continue
        xs, ys = [p[0] for p in pts], [p[1] for p in pts]
        f = up * scale
        out.append({"t": t, "x0": int(x0) / scale + min(xs) / f, "x1": int(x0) / scale + max(xs) / f, "y0": int(y0) / scale + min(ys) / f,
                    "y1": int(y0) / scale + max(ys) / f, "conf": float(score) * 100})
    return out


def _cells_from_words(ws: list[dict], box, lines: dict) -> tuple[list[list[str]], list[list[float]]]:
    """Each word into the cell (between the lines) its centre falls in; per cell its text and its lowest score. Rows with nothing are kept here."""
    x0, y0, x1, y1 = [float(v) for v in box]
    cx = [x0] + list(lines.get("cols", [])) + [x1]
    cy = [y0] + list(lines.get("rows", [])) + [y1]
    grid = [["" for _ in range(len(cx) - 1)] for _ in range(len(cy) - 1)]
    conf = [[100.0 for _ in range(len(cx) - 1)] for _ in range(len(cy) - 1)]
    for w in sorted(ws, key=lambda w: (round(w["y0"]), w["x0"])):
        mx, my = (w["x0"] + w["x1"]) / 2, (w["y0"] + w["y1"]) / 2
        c = next((i for i in range(len(cx) - 1) if cx[i] <= mx < cx[i + 1]), None)
        r = next((i for i in range(len(cy) - 1) if cy[i] <= my < cy[i + 1]), None)
        if c is None or r is None:
            continue
        grid[r][c] = (grid[r][c] + " " if grid[r][c] else "") + w["t"]
        conf[r][c] = min(conf[r][c], w["conf"])
    return grid, conf


def _ocr_words(info: dict, page: int, box: tuple[float, float, float, float]) -> list[dict]:
    """OCR on the box only: cut out, enlarged, greyscale, read as a block; words come back with their positions in the page's own coordinates."""
    import pytesseract
    from PIL import Image, ImageOps
    if info["ext"] == "pdf":
        scale = OCR_SCALE * 1.5
        img = _picture(info, page, scale)
        factor = scale
    else:
        img = _picture(info, page, 1.0)
        factor = 1.0
    x0, y0, x1, y1 = [v * factor for v in box]
    crop = img.crop((int(x0), int(y0), int(x1) + 1, int(y1) + 1))
    up = 1.0 if info["ext"] == "pdf" else OCR_SCALE
    if up != 1.0:
        crop = crop.resize((int(crop.width * up), int(crop.height * up)), Image.LANCZOS)
    crop = ImageOps.autocontrast(ImageOps.grayscale(crop))
    one_line = crop.height < crop.width * 0.3 and crop.height < 140 * OCR_SCALE          # a box around one figure: read as a single line (more accurate)
    data = pytesseract.image_to_data(crop, config="--psm 7" if one_line else "--psm 6", output_type=pytesseract.Output.DICT)
    out = []
    for i, t in enumerate(data["text"]):
        t = (t or "").strip()
        if not t:
            continue
        try:
            conf = float(data["conf"][i])
        except (TypeError, ValueError):
            conf = 0.0
        l, tp, w, h = data["left"][i], data["top"][i], data["width"][i], data["height"][i]
        out.append({"t": t, "x0": x0 / factor + l / up / factor, "x1": x0 / factor + (l + w) / up / factor,
                    "y0": y0 / factor + tp / up / factor, "y1": y0 / factor + (tp + h) / up / factor, "conf": conf})
    return out


_NOISE = re.compile(r"^[\W_]+$")


def _clean_cell(t: str) -> str:
    """OCR noise off a cell: a table border read as '|' or '[', stray dots and dashes; a cell that is only such marks is empty."""
    t = (t or "").strip()
    t = re.sub(r"^[|\[\]{}()!:;.,_~\-–—'`\"*]+", "", t)
    t = re.sub(r"[|\[\]{}()!:;_~–—'`\"*,.]+$", "", t)
    return "" if _NOISE.match(t or "x") else t.strip()


def _page_gray(info: dict, page: int, scale_pdf: float = 3.0):
    """The page as a greyscale picture plus the factor from the box's coordinates (PDF points, or picture pixels) to its pixels."""
    if info["ext"] == "pdf":
        return _picture(info, page, scale_pdf).convert("L"), scale_pdf
    return _picture(info, page, 1.0).convert("L"), 1.0


def _lines(profile, thresh: float, min_gap: int) -> list[float]:
    out, start = [], None
    for i, v in enumerate(profile):
        if v > thresh and start is None:
            start = i
        elif v <= thresh and start is not None:
            out.append((start + i - 1) / 2)
            start = None
    if start is not None:
        out.append((start + len(profile) - 1) / 2)
    merged: list[float] = []
    for q in out:
        if merged and q - merged[-1] < min_gap:
            merged[-1] = (merged[-1] + q) / 2
        else:
            merged.append(q)
    return merged


def table_lines(gray, box) -> tuple[list[float], list[float]]:
    """A ruled table in a picture: the horizontal and vertical lines inside the box, from runs of dark pixels (the box's own pixel coordinates)."""
    import numpy as np
    x0, y0, x1, y1 = [int(v) for v in box]
    a = np.asarray(gray.crop((x0, y0, x1, y1)))
    if a.size == 0:
        return [], []
    dark = a < 150
    h, w = dark.shape
    hl = _lines(dark.sum(axis=1) / max(1, w), 0.45, 8)
    if len(hl) >= 2:
        top, bot = int(hl[0]), int(hl[-1])
        vl = _lines(dark[top:bot + 1].sum(axis=0) / max(1, bot - top + 1), 0.5, 8)
    else:
        vl = _lines(dark.sum(axis=0) / max(1, h), 0.5, 8)
    return [y0 + v for v in hl], [x0 + v for v in vl]


def _ocr_one(gray, cell_box) -> tuple[str, float]:
    import pytesseract
    from PIL import Image, ImageOps
    x0, y0, x1, y1 = cell_box
    cell = gray.crop((x0, y0, x1, y1))
    cell = cell.resize((cell.width * OCR_SCALE, cell.height * OCR_SCALE), Image.LANCZOS)
    cell = ImageOps.autocontrast(cell).point(lambda v: 255 if v > 140 else 0)
    pad = Image.new("L", (cell.width + 20, cell.height + 20), 255)
    pad.paste(cell, (10, 10))
    d = pytesseract.image_to_data(pad, config="--psm 7", output_type=pytesseract.Output.DICT)
    parts, confs = [], []
    for t, cf in zip(d["text"], d["conf"]):
        t = _clean_cell(t or "")
        if t:
            parts.append(t)
            try:
                confs.append(float(cf))
            except (TypeError, ValueError):
                confs.append(0.0)
    return " ".join(parts), (min(confs) if confs else 100.0)


def _inner(edges: list[float], lo: float, hi: float, gap: float = 3.0) -> list[float]:
    """The lines strictly inside a box side (its own edges are not lines to move), sorted, none closer than `gap`."""
    out: list[float] = []
    for v in sorted(float(e) for e in edges):
        if lo + gap < v < hi - gap and (not out or v - out[-1] >= gap):
            out.append(round(v, 2))
    return out


def lines_from_words(ws: list[dict], box) -> dict:
    """Where the rows and columns of a table are, from its words: a row line halfway between two rows of words, a column line halfway between the
    end of one column's words and the start of the next's. {'cols': [x, ...], 'rows': [y, ...]} in the page's coordinates, inside the box."""
    x0, y0, x1, y1 = box
    rows = _rows(ws)
    ys = []
    for a, b in zip(rows, rows[1:]):
        bottom, top = max(w["y1"] for w in a), min(w["y0"] for w in b)
        ys.append((bottom + top) / 2 if top > bottom else top - 1)
    cells, cols = _cells_and_cols(ws)
    xs = []
    for k in range(1, len(cols)):
        left = [c["x1"] for r in cells for c in r if _col_of(c, cols) == k - 1]
        right = [c["x0"] for r in cells for c in r if _col_of(c, cols) == k]
        if left and right:
            le, ri = max(left), min(right)
            xs.append((le + ri) / 2 if ri > le else ri - 2)
        else:
            xs.append(cols[k]["x"] - 4)
    return {"cols": _inner(xs, x0, x1), "rows": _inner(ys, y0, y1)}


def _guess_lines(info: dict, page: int, box) -> dict:
    """Lines for a box without reading it: the words of a text PDF, else the ruled lines found in the picture (no OCR: fast)."""
    if info["ext"] == "pdf" and info["has_text"]:
        ws = [w for w in _pdf_words(info, page) if _inside(w, box)]
        return lines_from_words(ws, box) if ws else {"cols": [], "rows": []}
    gray, f = _page_gray(info, page)
    hl, vl = table_lines(gray, [v * f for v in box])
    return {"cols": _inner([v / f for v in vl], box[0], box[2]), "rows": _inner([v / f for v in hl], box[1], box[3])}


def merge_lines(mine: dict, guess: dict, near: float = 8.0) -> dict:
    """The person's lines win; a guessed line is added only where none of theirs is near (the box was made bigger: the new part gets the guess)."""
    out = {}
    for k in ("cols", "rows"):
        keep = [float(v) for v in mine.get(k, [])]
        out[k] = sorted(keep + [g for g in guess.get(k, []) if all(abs(g - m) > near for m in keep)])
    return out


def read_by_lines(info: dict, page: int, box, lines: dict) -> dict:
    """The person's own rows and columns: every cell between two lines is read on its own -- the words inside it (a text PDF, exact) or an OCR of
    just that cell (a picture). Rows with nothing in any cell are left out."""
    from concurrent.futures import ThreadPoolExecutor
    x0, y0, x1, y1 = [float(v) for v in box]
    cx = [x0] + _inner(lines.get("cols", []), x0, x1) + [x1]
    cy = [y0] + _inner(lines.get("rows", []), y0, y1) + [y1]
    used = {"cols": cx[1:-1], "rows": cy[1:-1]}
    cells = [[(cx[c], cy[r], cx[c + 1], cy[r + 1]) for c in range(len(cx) - 1)] for r in range(len(cy) - 1)]
    if info["ext"] == "pdf" and info["has_text"]:
        ws = [w for w in _pdf_words(info, page) if _inside(w, box)]
        grid = []
        for row in cells:
            line = []
            for (a, b, c, d) in row:
                inside = sorted((w for w in ws if _inside(w, (a, b, c, d))), key=lambda w: (round(w["y0"]), w["x0"]))
                line.append(" ".join(w["t"] for w in inside))
            grid.append(line)
        confs = [[100.0] * len(r) for r in grid]
    elif (rw := _rapid_words(info, page, box)) is not None:
        grid, confs = _cells_from_words(rw, box, used)
        confs = [[100.0 if c >= RAPID_DOUBT else 0.0 for c in r] for r in confs]
    else:
        gray, f = _page_gray(info, page)
        flat = [(int(a * f) + 2, int(b * f) + 2, int(c * f) - 1, int(d * f) - 1) for row in cells for (a, b, c, d) in row]
        with ThreadPoolExecutor(max_workers=4) as pool:
            got = list(pool.map(lambda cb: _ocr_one(gray, cb) if cb[2] - cb[0] >= 6 and cb[3] - cb[1] >= 6 else ("", 100.0), flat))
        w = len(cx) - 1
        grid = [[got[r * w + c][0] for c in range(w)] for r in range(len(cy) - 1)]
        confs = [[got[r * w + c][1] for c in range(w)] for r in range(len(cy) - 1)]
    keep = [i for i, r in enumerate(grid) if any(c.strip() for c in r)]
    return {"mode": "table", "rows": [grid[i] for i in keep], "doubt": [[c < OCR_DOUBT for c in confs[i]] for i in keep], "text": "",
            "how": "lines", "lines": used}


def _ocr_cells(info: dict, page: int, box) -> tuple[list[list[str]], list[list[float]], dict] | None:
    """A ruled table in a scan or photo: the table's own lines are found, each cell is cut out and read on its own (far more accurate than the block:
    JRCR's Yokohama photos, 2026-10-16), several cells at a time. None when the box holds no ruled table."""
    from concurrent.futures import ThreadPoolExecutor
    gray, f = _page_gray(info, page)
    pbox = [v * f for v in box]
    hl, vl = table_lines(gray, pbox)
    if len(hl) < 3 or len(vl) < 3:
        return None
    cells = []
    for r in range(len(hl) - 1):
        y0, y1 = int(hl[r]) + 2, int(hl[r + 1]) - 1
        if y1 - y0 < 6:
            continue
        row = []
        for c in range(len(vl) - 1):
            x0, x1 = int(vl[c]) + 2, int(vl[c + 1]) - 1
            row.append((x0, y0, x1, y1) if x1 - x0 >= 6 else None)
        cells.append(row)
    flat = [cb for row in cells for cb in row]
    with ThreadPoolExecutor(max_workers=4) as pool:
        got = list(pool.map(lambda cb: _ocr_one(gray, cb) if cb else ("", 100.0), flat))
    grid, confs, k = [], [], 0
    for row in cells:
        texts, cf = [], []
        for _ in row:
            t, c = got[k]
            k += 1
            texts.append(t)
            cf.append(c)
        if any(texts):
            grid.append(texts)
            confs.append(cf)
    if not grid:
        return None
    return grid, confs, {"cols": _inner([v / f for v in vl], box[0], box[2]), "rows": _inner([v / f for v in hl], box[1], box[3])}


def words(token: str, page: int) -> list[dict]:
    """The page's words with positions (a text PDF); [] for a picture (its words come only from an OCR of a box)."""
    info = file_info(token)
    if not info or not 1 <= page <= info["pages"] or info["ext"] != "pdf" or not info["has_text"]:
        return []
    return _pdf_words(info, page)


# --- reading a box ------------------------------------------------------------------------------------------------------------------------------------

def _inside(w: dict, box) -> bool:
    x0, y0, x1, y1 = box
    cx, cy = (w["x0"] + w["x1"]) / 2, (w["y0"] + w["y1"]) / 2
    return x0 <= cx <= x1 and y0 <= cy <= y1


def _rows(ws: list[dict]) -> list[list[dict]]:
    rows: list[dict] = []
    for w in sorted(ws, key=lambda w: (w["y0"], w["x0"])):
        cy, h = (w["y0"] + w["y1"]) / 2, (w["y1"] - w["y0"]) or 1
        r = next((r for r in rows if abs(r["y"] - cy) < h * 0.6), None)
        if r:
            r["w"].append(w)
        else:
            rows.append({"y": cy, "w": [w]})
    return [sorted(r["w"], key=lambda w: w["x0"]) for r in rows]


def _col_of(c: dict, cols: list[dict]) -> int:
    k = 0
    for i, col in enumerate(cols):
        if c["x0"] >= col["x"] - 18:
            k = i
    return k


def _grid(ws: list[dict]) -> tuple[list[list[str]], list[list[float]]]:
    """The position method: cells = words closer than a space and a half; column edges = where cells start across the rows. Returns the grid and,
    per cell, the lowest OCR confidence inside it (100 for a text PDF)."""
    cells, cols = _cells_and_cols(ws)
    grid, confs = [], []
    for r in cells:
        line, conf = [""] * len(cols), [100.0] * len(cols)
        for c in r:
            k = _col_of(c, cols)
            line[k] = (line[k] + " " if line[k] else "") + c["t"]
            conf[k] = min(conf[k], c["conf"])
        grid.append(line)
        confs.append(conf)
    return grid, confs


def _cells_and_cols(ws: list[dict]) -> tuple[list[list[dict]], list[dict]]:
    """Words joined into cells (closer than a space and a half), and the column starts found across the rows."""
    rows = _rows(ws)
    cells = []
    for r in rows:
        c: list[dict] = []
        for w in r:
            if c and w["x0"] - c[-1]["x1"] < 6:
                c[-1]["t"] += " " + w["t"]
                c[-1]["x1"] = w["x1"]
                c[-1]["conf"] = min(c[-1]["conf"], w["conf"])
            else:
                c.append(dict(w))
        cells.append(c)
    starts: list[dict] = []
    for r in cells:
        for c in r:
            s = next((s for s in starts if abs(s["x"] - c["x0"]) < 18), None)
            if s:
                s["n"] += 1
                s["x"] = (s["x"] * (s["n"] - 1) + c["x0"]) / s["n"]
            else:
                starts.append({"x": c["x0"], "n": 1})
    cols = sorted([s for s in starts if s["n"] >= max(2, len(cells) * 0.3)], key=lambda s: s["x"]) or [{"x": 0}]
    return cells, cols


def _real_cols(grid: list[list[str]]) -> int:
    """How many columns are filled in most rows (a stray label caught by the box is not a column)."""
    if not grid:
        return 0
    width = max(len(r) for r in grid)
    return sum(1 for i in range(width) if sum(1 for r in grid if i < len(r) and r[i].strip()) >= len(grid) * 0.6)


def _ruled_table(info: dict, page: int, box) -> list[list[str]] | None:
    """A table drawn with lines inside the box: pdfplumber's own reader on the crop; None when there is no such table."""
    import pdfplumber
    try:
        with pdfplumber.open(str(info["path"])) as pdf:
            p = pdf.pages[page - 1]
            x0, y0, x1, y1 = box
            m = 8                                                   # a border a few points outside the box still belongs to the table
            crop = p.crop((max(0, x0 - m), max(0, y0 - m), min(p.width, x1 + m), min(p.height, y1 + m)))
            if len(crop.lines) + len(crop.rects) < 4:
                return None
            table = crop.extract_table()
    except Exception:                                   # noqa: BLE001
        return None
    if not table or len(table) < 2:
        return None
    rows = [[" ".join(str(c or "").split()) for c in row] for row in table]
    rows = [r for r in rows if any(r)]
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [i for i in range(width) if any(r[i] for r in rows)]        # drop columns that are empty everywhere (a border read as a column)
    return [[r[i] for i in keep] for r in rows] if len(keep) >= 2 else None


def read_box(token: str, page: int, box, mode: str = "table", lines: dict | None = None, merge: bool = False) -> dict:
    """What is inside the box: {'mode', 'rows': [[cell, ...]], 'doubt': [[bool, ...]], 'text', 'how': 'ruled' | 'positions' | 'ocr' | 'lines' | 'none',
    'lines': {'cols', 'rows'}}. With `lines` (the person's rows and columns) each cell between them is read on its own; with `merge` too, the
    system's guess fills the part of the box the person's lines do not reach (a box made bigger)."""
    info = file_info(token)
    if not info or not 1 <= page <= info["pages"]:
        return {"mode": mode, "rows": [], "doubt": [], "text": "", "how": "none"}
    box = tuple(float(v) for v in box)
    if box[2] - box[0] < 4 or box[3] - box[1] < 4:
        return {"mode": mode, "rows": [], "doubt": [], "text": "", "how": "none"}
    if mode == "table" and lines is not None:
        mine = {k: [float(v) for v in (lines.get(k) or []) if isinstance(v, (int, float))] for k in ("cols", "rows")}
        if merge:
            mine = merge_lines(mine, _guess_lines(info, page, box))
        return read_by_lines(info, page, box, mine)
    r = _read_box(info, page, box, mode)
    if mode == "table" and r["how"] != "none" and "lines" not in r:
        r["lines"] = _guess_lines(info, page, box) if r["how"] == "ruled" else r.pop("_lines", {"cols": [], "rows": []})
    r.pop("_lines", None)
    return r


def _read_box(info: dict, page: int, box, mode: str) -> dict:
    how = "positions"
    if info["ext"] == "pdf" and info["has_text"]:
        ws = [w for w in _pdf_words(info, page) if _inside(w, box)]
        if mode == "table":
            ruled = _ruled_table(info, page, box)
            if ruled and ws:
                by_pos, _ = _grid(ws)
                if _real_cols(ruled) >= _real_cols(by_pos):               # a box that cuts a border loses a ruled column: then the positions win
                    return {"mode": mode, "rows": ruled, "doubt": [[False] * len(r) for r in ruled], "text": "", "how": "ruled",
                            "lines": lines_from_words(ws, box)}
    else:
        rw = _rapid_words(info, page, box)
        if rw is not None:
            if not rw:
                return {"mode": mode, "rows": [], "doubt": [], "text": "", "how": "none"}
            if mode == "text":
                return {"mode": mode, "rows": [], "doubt": [], "text": " ".join(" ".join(w["t"] for w in r) for r in _rows(rw)), "how": "ocr",
                        "doubtful": any(w["conf"] < RAPID_DOUBT for w in rw)}
            ruled = _guess_lines(info, page, box)
            if len(ruled["cols"]) >= 1 and len(ruled["rows"]) >= 1:
                grid, confs = _cells_from_words(rw, box, ruled)
                ruled = {"cols": list(ruled["cols"]), "rows": ruled["rows"]}
                while ruled["cols"] and grid and not any(r[0].strip() for r in grid):          # the table's own outer border inside the box: no empty column
                    grid, confs, ruled["cols"] = [r[1:] for r in grid], [r[1:] for r in confs], ruled["cols"][1:]
                while ruled["cols"] and grid and not any(r[-1].strip() for r in grid):
                    grid, confs, ruled["cols"] = [r[:-1] for r in grid], [r[:-1] for r in confs], ruled["cols"][:-1]
                keep = [i for i, r in enumerate(grid) if any(c.strip() for c in r)]
                return {"mode": mode, "rows": [grid[i] for i in keep], "doubt": [[c < RAPID_DOUBT for c in confs[i]] for i in keep], "text": "",
                        "how": "ocr", "lines": ruled}
            grid, confs = _grid(rw)
            return {"mode": mode, "rows": grid, "doubt": [[c < RAPID_DOUBT for c in r] for r in confs], "text": "", "how": "ocr", "lines": lines_from_words(rw, box)}
        if mode == "table":
            cells = _ocr_cells(info, page, box)
            if cells:
                grid, confs, found = cells
                return {"mode": mode, "rows": grid, "doubt": [[c < OCR_DOUBT for c in r] for r in confs], "text": "", "how": "ocr", "lines": found}
        ws = [w for w in _ocr_words(info, page, box) if _clean_cell(w["t"])]
        for w in ws:
            w["t"] = _clean_cell(w["t"])
        how = "ocr"
    if not ws:
        return {"mode": mode, "rows": [], "doubt": [], "text": "", "how": "none"}
    if mode == "text":
        text = " ".join(" ".join(w["t"] for w in r) for r in _rows(ws))
        return {"mode": mode, "rows": [], "doubt": [], "text": text, "how": how, "doubtful": any(w["conf"] < OCR_DOUBT for w in ws)}
    grid, confs = _grid(ws)
    return {"mode": mode, "rows": grid, "doubt": [[c < OCR_DOUBT for c in r] for r in confs], "text": "", "how": how, "lines": lines_from_words(ws, box)}


# --- remembering a snip while its rows travel to the import (deleted with the document) ------------------------------------------------------------------

def save_snip(token: str, page: int, box, mode: str, rows: list, used_for: str, ref_id: int | None = None) -> int | None:
    info = file_info(token)
    if not info:
        return None
    with write_tx() as conn:
        return conn.execute("""INSERT INTO snips (file_id, page, x0, y0, x1, y1, mode, rows, used_for, ref_id, created_by, created_at)
                               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, acting_user(), ?) RETURNING snip_id""",
                            (info["file_id"], int(page), *[round(float(v), 2) for v in box], mode, json.dumps(rows)[:200000], used_for, ref_id,
                             manila_now().strftime(STAMP))).fetchone()[0]


def snip_info(snip_id: int) -> dict | None:
    with get_conn() as conn:
        row = conn.execute("""SELECT s.snip_id, f.token, f.filename, f.kind, s.page, s.x0, s.y0, s.x1, s.y1, s.mode, s.used_for, s.ref_id, s.created_by, s.created_at
                              FROM snips s JOIN snip_files f ON f.file_id = s.file_id WHERE s.snip_id = ?""", (int(snip_id),)).fetchone()
    if not row:
        return None
    keys = ("snip_id", "token", "filename", "kind", "page", "x0", "y0", "x1", "y1", "mode", "used_for", "ref_id", "by", "at")
    return dict(zip(keys, row))


# --- the price list use (stage 1) -----------------------------------------------------------------------------------------------------------------------

_PRICE = re.compile(r"^[₱P]?\s*\d{1,3}(,\d{3})+(\.\d{1,2})?$|^[₱P]?\s*\d+(\.\d{1,2})?$")
_DATE = re.compile(r"^\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}$|^\d{1,2}[ -][A-Za-z]{3,9}[ -]\d{2,4}$|^[A-Za-z]{3,9}\.? \d{1,2},? \d{4}$|^\d{4}-\d{2}-\d{2}$")
_PLY = re.compile(r"^\d{1,2}\s*PR$|^\d{2,3}[A-Z]$", re.I)


def has_header(rows: list[list[str]]) -> bool:
    """The first row is a heading when none of its cells looks like a price and some cell is a known column word."""
    if not rows:
        return False
    first = [c.strip() for c in rows[0] if c.strip()]
    if not first or any(_PRICE.match(c) or _DATE.match(c) for c in first):
        return False
    words = {"size", "sizes", "pattern", "price", "model", "brand", "list price", "srp", "description", "item", "ply", "terrain", "type", "due date", "amount",
             "interest", "principal", "balance", "date", "no.", "amortization", "remarks", "origin", "rim", "tire size"}
    lettered = sum(1 for c in first if any(ch.isalpha() for ch in c))
    return any(c.lower() in words for c in first) or lettered >= max(2, len(first) * 0.7)


def guess_colmap(category: str, rows: list[list[str]]) -> list[str]:
    """Owner, 2026-10-16: "since we have fields already in place ... it filters the word it will snip". Names each snipped column from what is in it,
    with the category's own fields: a column of prices is the List price, of tire sizes the Size, of Lug/Rib/Mix the Terrain, of 20PR the Ply, a
    heading word is taken at its word, the rest is the Pattern / Model. The person corrects a guess on the screen."""
    from engines import product_catalog as pc
    from engines import price_template_service as pt
    if not rows:
        return []
    width = max(len(r) for r in rows)
    keys = [k for k, _, _ in pt._headers(category)]
    labels = {label.lower(): k for k, label, _ in pt._headers(category)}
    labels.update({"sizes": "size", "price": "list_price", "srp": "list_price", "list": "list_price", "model": "pattern" if category == "Tires" else "model",
                   "pattern": "pattern", "description": "pattern" if category == "Tires" else "model", "item": "pattern" if category == "Tires" else "model",
                   "code": "supplier_code", "item code": "supplier_code", "date": "effective_date", "effective": "effective_date", "remarks": "remarks", "remark": "remarks"})
    head = [c.strip().lower() for c in rows[0]] + [""] * width if has_header(rows) else [""] * width
    body = rows[1:] if has_header(rows) else rows
    out, used = [], set()
    for i in range(width):
        col = [r[i].strip() for r in body if i < len(r) and r[i].strip()]
        guess = ""
        hit = labels.get(head[i]) or next((v for k, v in labels.items() if head[i] and (k in head[i] or head[i] in k) and len(k) > 3), "")
        if hit and hit in keys:
            guess = hit
        elif col:
            share = lambda f: sum(1 for c in col if f(c)) / len(col)          # noqa: E731
            if share(lambda c: bool(_PRICE.match(c))) >= 0.7:
                guess = "list_price"
            elif category == "Tires" and share(lambda c: bool(pc.TIRE_SIZE.search(c.replace(" ", "").upper()))) >= 0.6:
                guess = "size"
            elif category == "Tires" and share(lambda c: pc.terrain_of(c) is not None) >= 0.7:
                guess = "terrain"
            elif category == "Tires" and share(lambda c: bool(_PLY.match(c))) >= 0.7:
                guess = "ply"
            elif share(lambda c: bool(re.fullmatch(r"\d{1,2}/\d{1,2}/\d{2,4}", c))) >= 0.7:
                guess = "effective_date"
            elif share(lambda c: bool(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9 /.+-]{1,24}", c))) >= 0.7:
                guess = "pattern" if category == "Tires" else ("model" if category != "Rims" else "size")
        if guess in used and guess not in ("remarks",):
            guess = "remarks" if "remarks" not in used else ""
        out.append(guess if guess in keys else "")
        if guess:
            used.add(guess)
    return out



def price_workbook(creditor_id: int, category: str, brand: str, rows: list[list[str]], colmap: list[str], start, header_row: bool = True) -> bytes:
    """The snipped rows as the category's own template sheet (the same checks and product key as an Excel import: price_template_service.parse_workbook).
    colmap: per snipped column, the template key it feeds ('' = not used); brand and the start date fill every row unless a mapped column gives them."""
    from openpyxl import Workbook

    from engines import price_template_service as pt
    wb = Workbook()
    guide = wb.active
    guide.title = "How to fill"
    guide["A1"] = pt.TYPE_HOW_TO[0]
    guide["A2"] = pt.SUPPLIER_LINE + pt._supplier_name(creditor_id)
    ws = wb.create_sheet(category)
    heads = pt._headers(category)
    ws.append([pt._title(label, req) for _, label, req in heads])
    body = rows[1:] if header_row and len(rows) > 1 else rows
    for r in body:
        line = {}
        for i, key in enumerate(colmap):
            if key and i < len(r) and r[i].strip():
                line[key] = (line[key] + " " if line.get(key) else "") + r[i].strip()
        if not any(v for k, v in line.items() if k not in ("brand",)):
            continue
        if category == "Tires" and line.get("size"):                 # the size cell as the supplier prints it: size, ply, terrain, the rest to Remarks
            bits = split_size_cell(line["size"])
            if bits["size"]:
                line["size"] = bits["size"]
                for k in ("ply", "terrain"):
                    if bits[k] and not line.get(k):
                        line[k] = bits[k]
                if bits["rest"]:
                    line["remarks"] = (line.get("remarks", "") + " " + bits["rest"]).strip()
        line.setdefault("brand", brand)
        if not line.get("effective_date") and start:
            line["effective_date"] = start
        ws.append([line.get(key, "") or "" for key, _, _ in heads])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def split_size_cell(text: str) -> dict:
    """A supplier's size cell as printed ("315/80R22.5-20PR Lug", "12.00R20-18ply MIX", "23.5-25-20PR (TUBETYPE)") -> {'size', 'ply', 'terrain',
    'rest'}: the size the catalog knows, the ply as 20PR, the terrain word, and what else was there (to Remarks). Owner, 2026-10-16: "it filters the
    word it will snip"."""
    from engines import product_catalog as pc
    t = " ".join((text or "").split())
    out = {"size": "", "ply": "", "terrain": "", "rest": ""}
    m = pc.TIRE_SIZE.search(t.replace(" ", "").upper())
    if not m:
        return {**out, "rest": t}
    out["size"] = m.group(0)
    rest = t
    # take the size out of the text as printed (it may have been written with spaces), then the ply, then the terrain
    rest = re.sub(r"\s*" + r"\s*".join(re.escape(ch) for ch in m.group(0)) + r"\s*", " ", rest, count=1, flags=re.I)
    if p := re.search(r"-?\s*(\d{1,2})\s*(PR|PLY)\b\.?", rest, flags=re.I):
        out["ply"] = p.group(1) + "PR"
        rest = rest[:p.start()] + " " + rest[p.end():]
    elif p := re.match(r"\s*-\s*(\d{1,2})\b", rest):                     # "12.00R20-18 Lug": a truck tire's ply written without PR
        if 4 <= int(p.group(1)) <= 30:
            out["ply"] = p.group(1) + "PR"
            rest = rest[p.end():]
    words = []
    for w in rest.replace("-", " ").split():
        terr = pc.terrain_of(w)
        if terr and not out["terrain"]:
            out["terrain"] = terr
        else:
            words.append(w)
    out["rest"] = " ".join(words).strip(" -,;")
    return out




# --- the bank loan use (stage 2; owner, 2026-10-16: "I can use it as well for bank loans when they give amortization table") ----------------------------

LOAN_KEYS = [("due_date", "Due date"), ("amount", "Amount (the amortization)"), ("interest", "Interest part"), ("principal", "Principal (not used)"), ("balance", "Balance (not used)"), ("no", "No. (not used)")]


def _loan_header(word: str) -> str:
    w = word.lower()
    if "interest" in w:
        return "interest"
    if "principal" in w:
        return "principal"
    if "balance" in w or "outstanding" in w:
        return "balance"
    if "due" in w or "date" in w:
        return "due_date"
    if any(k in w for k in ("amort", "total", "payment", "installment", "instalment", "amount")):
        return "amount"
    if w in ("no", "no.", "#", "n", "mo", "month", "period", "seq"):
        return "no"
    return ""


def guess_loan_colmap(rows: list[list[str]]) -> list[str]:
    """The amortisation table's columns: Due date, Amount, Interest part; Principal, Balance and No. are shown but not used. Headings first; without
    them, the dates are the due date and the first money column is the amount."""
    if not rows:
        return []
    width = max(len(r) for r in rows)
    head = has_header(rows) or any(_loan_header(c) for c in rows[0])
    body = rows[1:] if head else rows
    out, used = [], set()
    for i in range(width):
        col = [r[i].strip() for r in body if i < len(r) and r[i].strip()]
        guess = _loan_header(rows[0][i]) if head and i < len(rows[0]) else ""
        if not guess and col:
            share = lambda f: sum(1 for c in col if f(c)) / len(col)          # noqa: E731
            if share(lambda c: bool(_DATE.match(c))) >= 0.7:
                guess = "due_date"
            elif share(lambda c: bool(_PRICE.match(c))) >= 0.7:
                guess = "amount" if "amount" not in used else ("interest" if "interest" not in used else "balance")
            elif share(lambda c: c.isdigit()) >= 0.7:
                guess = "no"
        if guess in used and guess != "":
            guess = ""
        out.append(guess)
        if guess:
            used.add(guess)
    return out


def _money(text: str) -> float | None:
    t = (text or "").replace("₱", "").replace("PHP", "").replace(",", "").replace(" ", "").strip()
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    try:
        v = float(t)
    except ValueError:
        return None
    return -v if neg else v


def loan_rows(rows: list[list[str]], colmap: list[str], header_row: bool) -> tuple[list[dict], list[str]]:
    """The snipped table as schedule rows: ([{'due': 'YYYY-MM-DD', 'amount', 'interest'}], problems). A row that cannot be read is a problem, not a
    silent drop (the person types it)."""
    import pandas as pd
    want = {k: i for i, k in enumerate(colmap) if k}
    out, problems = [], []
    if "due_date" not in want or "amount" not in want:
        return [], ["Name the Due date and the Amount columns."]
    body = rows[1:] if header_row and len(rows) > 1 else rows
    for n, r in enumerate(body, start=2 if header_row else 1):
        cell = lambda k: r[want[k]].strip() if k in want and want[k] < len(r) else ""       # noqa: E731
        if not any(x.strip() for x in r):
            continue
        try:
            due = pd.to_datetime(cell("due_date"), dayfirst=False).date()
        except (ValueError, TypeError):
            problems.append(f"Row {n}: the due date '{cell('due_date')}' could not be read.")
            continue
        amount = _money(cell("amount"))
        if amount is None or amount <= 0:
            problems.append(f"Row {n}: the amount '{cell('amount')}' is not a number above zero.")
            continue
        interest = _money(cell("interest")) if "interest" in want else 0.0
        if interest is None:
            problems.append(f"Row {n}: the interest '{cell('interest')}' is not a number.")
            continue
        out.append({"due": due.isoformat(), "amount": round(amount, 2), "interest": round(max(interest, 0.0), 2)})
    return out, problems


# --- slow reads in the background (2026-10-16: a photographed page takes a minute or two of OCR on the server's one CPU; a web request may not wait
# that long, so the read runs in a thread and writes its result to a file the page polls; files, so every gunicorn worker sees the same job) ------------

JOB_MINUTES = 6


def needs_job(token: str) -> bool:
    """A picture, or a PDF page without real text: read by OCR in the background (RapidOCR takes seconds; Tesseract, the fallback, minutes)."""
    info = file_info(token)
    return bool(info) and not (info["ext"] == "pdf" and info["has_text"])


def _jobs() -> Path:
    p = folder() / "jobs"
    p.mkdir(parents=True, exist_ok=True)
    for old in p.glob("*.json"):                                 # a day-old result is of no use to anyone
        try:
            if time.time() - old.stat().st_mtime > 86400:
                old.unlink(missing_ok=True)
        except OSError:
            pass
    return p


def start_job(token: str, page: int, box, mode: str, lines: dict | None = None, merge: bool = False) -> str:
    """Starts the read in the background; returns the job's id. The result lands in jobs/<id>.json."""
    import threading
    job = secrets.token_hex(8)
    path = _jobs() / f"{job}.json"
    path.write_text(json.dumps({"state": "running", "started": time.time()}), encoding="utf-8")

    def run():
        try:
            r = read_box(token, page, box, mode, lines, merge)
            r["state"] = "done"
        except Exception as e:                                   # noqa: BLE001 -- the OCR program missing, a broken file: the page says so
            r = {"state": "failed", "msg": f"Could not read the box: {e}"}
        path.write_text(json.dumps(r), encoding="utf-8")
    threading.Thread(target=run, daemon=True).start()
    return job


def job_result(job: str) -> dict | None:
    """{'state': 'running' | 'done' | 'failed' | 'lost', ...}; None for an unknown id."""
    if not job or not re.fullmatch(r"[0-9a-f]{16}", job):
        return None
    path = _jobs() / f"{job}.json"
    if not path.exists():
        return None
    try:
        r = json.loads(path.read_text(encoding="utf-8"))
    except ValueError:
        return {"state": "running"}
    if r.get("state") == "running" and time.time() - float(r.get("started") or 0) > JOB_MINUTES * 60:
        return {"state": "lost", "msg": f"The read took more than {JOB_MINUTES} minutes and was given up. Draw a smaller box, or try a clearer picture."}
    return r


# --- the vendor SOA use (stage 3; owner, 2026-10-16) ------------------------------------------------------------------------------------------------------

SOA_KEYS = [("ref", "SI / document no."), ("amount", "Amount"), ("date", "Date (not used)"), ("balance", "Balance (not used)"), ("other", "Other (not used)")]


def _soa_header(word: str) -> str:
    w = word.lower()
    if "balance" in w or "outstanding" in w:
        return "balance"
    if any(k in w for k in ("amount", "debit", "charge", "total", "value")):
        return "amount"
    if any(k in w for k in ("si", "invoice", "inv", "ref", "doc", "no.", "number", "dr")):
        return "ref"
    if "date" in w:
        return "date"
    return ""


def guess_soa_colmap(rows: list[list[str]]) -> list[str]:
    """The statement's columns: the SI / document no. and the amount; a date and a running balance are shown but not used. Headings first; without
    them, dates are dates, numbers of 4+ digits are the reference, the first money column is the amount and the next the balance."""
    if not rows:
        return []
    width = max(len(r) for r in rows)
    head = has_header(rows) or any(_soa_header(c) for c in rows[0])
    body = rows[1:] if head else rows
    out, used = [], set()
    for i in range(width):
        col = [r[i].strip() for r in body if i < len(r) and r[i].strip()]
        guess = _soa_header(rows[0][i]) if head and i < len(rows[0]) else ""
        if not guess and col:
            share = lambda f: sum(1 for c in col if f(c)) / len(col)          # noqa: E731
            if share(lambda c: bool(_DATE.match(c))) >= 0.7:
                guess = "date"
            elif share(lambda c: bool(_PRICE.match(c)) and ("," in c or "." in c)) >= 0.7:     # money is written with a comma or centavos
                guess = "amount" if "amount" not in used else "balance"
            elif share(lambda c: bool(re.fullmatch(r"[A-Za-z]{0,4}[-/ ]?\d{4,}[A-Za-z0-9/-]*", c))) >= 0.6:
                guess = "ref"                                                               # 244911, SI-0077: a document number
            elif share(lambda c: bool(_PRICE.match(c))) >= 0.7:
                guess = "amount" if "amount" not in used else "balance"
        if guess in used and guess:
            guess = ""
        out.append(guess)
        if guess:
            used.add(guess)
    return out


def soa_rows(rows: list[list[str]], colmap: list[str], header_row: bool) -> tuple[list[dict], list[str]]:
    """The snipped table as statement lines: ([{'ref', 'amount'}], problems). A line whose amount cannot be read is kept with amount 0 and said, so it
    shows as a difference rather than vanishing."""
    want = {k: i for i, k in enumerate(colmap) if k}
    if "ref" not in want or "amount" not in want:
        return [], ["Name the SI / document no. and the Amount columns."]
    body = rows[1:] if header_row and len(rows) > 1 else rows
    out, problems = [], []
    for n, r in enumerate(body, start=2 if header_row else 1):
        cell = lambda k: r[want[k]].strip() if want[k] < len(r) else ""         # noqa: E731
        ref = " ".join(cell("ref").split())
        if not ref:
            if any(x.strip() for x in r):
                problems.append(f"Row {n}: no document number (a total line?): left out.")
            continue
        amount = _money(cell("amount"))
        if amount is None:
            problems.append(f"Row {n} ({ref}): the amount '{cell('amount')}' could not be read: counted as 0.")
            amount = 0.0
        out.append({"ref": ref, "amount": round(amount, 2)})
    return out, problems

