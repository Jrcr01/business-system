"""Downloadable Excel invoice-entry workbook (macro-enabled) with today's supplier list.

The master file (excel_invoice_entry/invoice_entry_master.xlsm) carries the macros and the pop-up
form. On every download this loads it, drops any batch tabs left over from testing, rewrites the
_Suppliers tab from the system's active suppliers, and returns the result -- so the supplier list is
always as current as the moment of download. The VBA project is carried through untouched.
If the form is ever improved, replace the master file once; nothing here needs to change.
"""
import io
import os
import sqlite3
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from engines.common import get_conn

DB_PATH = os.environ.get("AP_DB_PATH", "data/ap_disbursement.db")
MASTER_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "excel_invoice_entry", "invoice_entry_master.xlsm")
DEFAULT_TERMS = "120"


def active_suppliers() -> list[tuple[str, str]]:
    with get_conn(DB_PATH) as conn:
        return conn.execute(
            "SELECT creditor_name, COALESCE(NULLIF(TRIM(terms), ''), ?) FROM creditors "
            "WHERE creditor_type = 'SUPPLIER' AND is_active = 1 ORDER BY LOWER(creditor_name)",
            (DEFAULT_TERMS,),
        ).fetchall()


def build_invoice_entry_workbook() -> bytes:
    suppliers = active_suppliers()
    wb = load_workbook(MASTER_PATH, keep_vba=True)
    for name in list(wb.sheetnames):
        if not name.startswith("_"):
            del wb[name]  # leftover batch tabs from testing; the form recreates tabs as needed
    if "_Suppliers" in wb.sheetnames:
        del wb["_Suppliers"]
    ws = wb.create_sheet("_Suppliers")
    ws.append(["Supplier", "Terms (days)"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for row in suppliers:
        ws.append(list(row))
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 14
    if "_Start" in wb.sheetnames:
        wb["_Start"]["B12"] = f"Supplier list refreshed from the AP system on {datetime.now():%Y-%m-%d %H:%M} ({len(suppliers)} active suppliers)."
        wb["_Start"]["B12"].font = Font(italic=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
