"""Friendly column names and CSV / Excel files for every report (2026-09-28, Reports review, plan phase 3). One list of names so the same
column reads the same on every screen and in every export: "check_date" -> "Check date", "creditor_name" -> "Payee". Money stays a number
(two decimals, thousands separators in Excel) and dates stay dates, so the files can be added up and sorted in a spreadsheet.
Screen helpers built on this: views/shared.py `show_table`, `export_buttons`."""
from __future__ import annotations

import io

import pandas as pd

COLUMN_LABELS = {
    "amount": "Amount", "total_amount": "Total", "gross_amount": "Gross", "net_total": "Net total", "paid": "Paid", "scheduled": "Post-dated",
    "open": "Not yet covered", "balance": "Still owed", "running_balance": "Running balance", "total_amount_due": "Amount due",
    "principal_amount": "Principal", "total_interest": "Interest", "total_payable": "Total payable", "outstanding_uncleared": "Outstanding",
    "outstanding": "Outstanding", "total_unpaid": "Total unpaid", "current_due": "Current", "days_1_30": "1-30 days", "days_31_60": "31-60 days",
    "days_over_60": "60+ days", "credit_limit": "Credit limit", "debit": "Out", "credit": "In", "unpaid_balance_next": "Carried over",
    "amount_paid": "Amount paid", "previous_balance": "Previous balance", "new_charges": "New charges", "interest_charged": "Interest charged",
    "interest": "Interest", "interest_portion": "Interest part", "our_amount": "Our amount",
    "check_date": "Check date", "due_date": "Due date", "actual_payment_date": "Payment date", "invoice_date": "Invoice date",
    "issue_date": "Issue date", "start_date": "Start date", "maturity_date": "Maturity date", "payment_date": "Payment date",
    "first_check": "First check", "last_check": "Last check", "posting_date": "Bank date", "bank_cleared_date": "Bank cleared",
    "created": "Created", "date_received": "Received", "received_date": "Received", "period_from": "From", "period_to": "To",
    "created_at": "Created", "updated_at": "Updated", "printed_at": "Printed", "imported_at": "Imported", "timestamp": "When",
    "check_number": "Check #", "creditor_name": "Payee", "payee_name": "Payee", "payee": "Payee", "supplier_name": "Supplier",
    "supplier": "Supplier", "lender_name": "Lender", "bank_name": "Bank", "statement_ref": "Batch", "batch": "Batch", "note_ref": "Loan note",
    "loan_ref": "Loan ref", "doc_ref": "Invoice #", "invoice": "Invoice", "reference": "Reference", "track": "Track", "status": "Status",
    "open_invoices": "Open invoices", "num_items": "# Items", "num_checks": "# Checks", "card_name": "Card", "card_number_last4": "Last 4",
    "issuer": "Issuer", "cardholder": "Cardholder", "installment_no": "#", "description": "Description", "kind": "Kind", "note": "Note",
    "resolution": "How resolved", "match_status": "Match", "user_action": "Action", "entity_name": "Record type", "entity_id": "Record id",
    "old_values": "Before", "new_values": "After", "monthly_interest_rate": "Monthly rate", "terms": "Terms", "days": "Days",
}
MONEY_COLUMNS = {
    "amount", "total_amount", "gross_amount", "net_total", "paid", "scheduled", "open", "balance", "running_balance", "total_amount_due",
    "principal_amount", "total_interest", "total_payable", "outstanding_uncleared", "outstanding", "total_unpaid", "current_due", "days_1_30",
    "days_31_60", "days_over_60", "credit_limit", "debit", "credit", "unpaid_balance_next", "amount_paid", "previous_balance", "new_charges",
    "interest_charged", "interest", "interest_portion", "our_amount",
}
DATE_COLUMNS = {
    "check_date", "due_date", "actual_payment_date", "invoice_date", "issue_date", "start_date", "maturity_date", "payment_date", "first_check",
    "last_check", "posting_date", "bank_cleared_date", "created", "date_received", "received_date", "period_from", "period_to",
}
DATETIME_COLUMNS = {"created_at", "updated_at", "printed_at", "imported_at", "timestamp"}


def label(col: str, labels: dict | None = None) -> str:
    """The friendly name of a column: the caller's own name first, then the shared list, then 'some_name' -> 'Some name'."""
    if labels and col in labels:
        return labels[col]
    if col in COLUMN_LABELS:
        return COLUMN_LABELS[col]
    if "_" not in str(col) and str(col) != str(col).lower():          # already a friendly name ("Sum of Amount", "Total")
        return str(col)
    return str(col).replace("_", " ").strip().capitalize()


def kind_of(col: str, kinds: dict | None = None) -> str:
    """'money' / 'date' / 'datetime' / 'number' / 'other' -- the caller's own `kinds` first (the report builder names its columns itself),
    then by name (the shared lists), so a money column is formatted the same everywhere."""
    if kinds and col in kinds:
        return kinds[col]
    if col in MONEY_COLUMNS:
        return "money"
    if col in DATE_COLUMNS:
        return "date"
    if col in DATETIME_COLUMNS:
        return "datetime"
    return "other"


def with_real_dates(df: pd.DataFrame, kinds: dict | None = None) -> pd.DataFrame:
    """Date and date-time columns (stored as text in SQLite) turned into real dates, for date formatting on screen and in Excel."""
    out = df.copy()
    for c in out.columns:
        k = kind_of(c, kinds)
        if k == "date":
            out[c] = pd.to_datetime(out[c], errors="coerce").dt.date
        elif k == "datetime":
            out[c] = pd.to_datetime(out[c], errors="coerce")
    return out


def friendly(df: pd.DataFrame, labels: dict | None = None) -> pd.DataFrame:
    """The same rows with friendly column names (for CSV / Excel)."""
    return df.rename(columns={c: label(c, labels) for c in df.columns})


def csv_bytes(df: pd.DataFrame, labels: dict | None = None, kinds: dict | None = None) -> bytes:
    """CSV with friendly headers; money as plain numbers with two decimals, dates as YYYY-MM-DD (spreadsheet-friendly). UTF-8 with a BOM
    so Excel shows ₱ and accents correctly when the file is double-clicked."""
    out = with_real_dates(df, kinds)
    for c in out.columns:
        if kind_of(c, kinds) == "money":
            out[c] = pd.to_numeric(out[c], errors="coerce").round(2)
    return friendly(out, labels).to_csv(index=False, float_format="%.2f").encode("utf-8-sig")


def excel_bytes(sheets: dict[str, pd.DataFrame] | pd.DataFrame, labels: dict | None = None, sheet_name: str = "Report",
                kinds: dict | None = None) -> bytes:
    """An .xlsx workbook, one sheet per table: friendly headers in bold, money as #,##0.00, dates as Mmm dd, yyyy, columns sized to fit,
    the header row frozen."""
    from openpyxl.styles import Font

    if isinstance(sheets, pd.DataFrame):
        sheets = {sheet_name: sheets}
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            safe = "".join(ch for ch in str(name) if ch not in '[]:*?/\\')[:31] or "Sheet"
            data = with_real_dates(df, kinds)
            friendly(data, labels).to_excel(xw, sheet_name=safe, index=False)
            ws = xw.sheets[safe]
            ws.freeze_panes = "A2"
            for i, col in enumerate(data.columns, start=1):
                ws.cell(row=1, column=i).font = Font(bold=True)
                k = kind_of(col, kinds)
                fmt = {"money": "#,##0.00", "number": "#,##0.00", "count": "#,##0", "date": "mmm dd, yyyy", "datetime": "mmm dd, yyyy h:mm AM/PM"}.get(k)
                if fmt:
                    for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                        row[0].number_format = fmt
                longest = max([len(label(col, labels))] + [len(f"{v:,.2f}" if k == "money" and isinstance(v, (int, float)) else str(v))
                                                             for v in data[col].head(500).tolist()])
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(10, longest + 2), 60)
    return buf.getvalue()
