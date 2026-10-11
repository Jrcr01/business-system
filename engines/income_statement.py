"""The month's Income Statement (owner, 2026-10-11: "just focus on income statement" -- no GL codes, no debits, nothing to maintain). Every figure
is computed straight from the documents: sales net of returns and voids, service income, cost of goods at the moving average the stock ledger keeps,
the owner's expense categories, shrinkage from approved counts and admin adjustments, interest actually paid to lenders (the installment's own
interest, when its check cleared) and on bank loans (the interest portion typed off the bank's document, when the installment was ticked paid).
Accrual by nature: a charge sale counts in the month it was sold."""
from __future__ import annotations

from engines.clearing import bank_loan_cleared_sql, cleared_sql
from engines.common import get_conn

# the month a payment counts in: the day it cleared (the bank's date, else the tick, else settled outside the bank)
_CHECK_DAY = "COALESCE(c.bank_cleared_date, c.manual_cleared_at, c.settled_outside_at)"
_LOAN_DAY = "COALESCE(bli.cleared_at, (SELECT MIN(t.txn_date) FROM bank_transactions t WHERE t.matched_installment_id = bli.installment_id))"


def _one(conn, sql: str, params=()) -> float:
    return round(float(conn.execute(sql, params).fetchone()[0] or 0), 2)


def month_report(month: str) -> dict:
    """month = '2026-10'. Lines for the page; every number traceable to its documents."""
    with get_conn() as conn:
        sales_goods = _one(conn, """SELECT COALESCE(SUM(l.amount), 0) FROM sale_lines l JOIN sales s ON s.sale_id = l.sale_id
                                    WHERE s.status = 'SAVED' AND s.is_opening = 0 AND l.item_id IS NOT NULL AND substr(s.sale_date, 1, 7) = ?""", (month,))
        services = _one(conn, """SELECT COALESCE(SUM(l.amount), 0) FROM sale_lines l JOIN sales s ON s.sale_id = l.sale_id
                                 WHERE s.status = 'SAVED' AND s.is_opening = 0 AND l.item_id IS NULL AND substr(s.sale_date, 1, 7) = ?""", (month,))
        returns = _one(conn, "SELECT COALESCE(SUM(amount), 0) FROM sale_returns WHERE substr(returned_on, 1, 7) = ?", (month,))
        cogs = _one(conn, """SELECT COALESCE(SUM(s.cost_total), 0) FROM sales s
                             WHERE s.status = 'SAVED' AND s.is_opening = 0 AND substr(s.sale_date, 1, 7) = ?""", (month,))
        returns_cost = _one(conn, """SELECT COALESCE(SUM(r.qty * l.unit_cost), 0) FROM sale_returns r JOIN sale_lines l ON l.line_id = r.line_id
                                     WHERE substr(r.returned_on, 1, 7) = ?""", (month,))
        shrinkage = -_one(conn, """SELECT COALESCE(SUM(qty * unit_cost), 0) FROM stock_ledger
                                   WHERE kind = 'ADJ' AND ref_type IN ('stock_count_lines', 'items') AND substr(moved_on, 1, 7) = ?""", (month,))
        # paid = cleared, by the one rule every report uses (engines/clearing.py): a cleared check keeps status 'Active' and carries its mark
        lender_interest = _one(conn, f"""SELECT COALESCE(SUM(i.interest), 0) FROM loan_installments i JOIN checks_issued c ON c.check_id = i.check_id
                                         WHERE c.status = 'Active' AND {cleared_sql('c')} AND substr({_CHECK_DAY}, 1, 7) = ?""", (month,))
        bankloan_interest = _one(conn, f"""SELECT COALESCE(SUM(bli.interest_portion), 0) FROM bank_loan_installments bli
                                           WHERE {bank_loan_cleared_sql('bli')} AND substr({_LOAN_DAY}, 1, 7) = ?""", (month,))
        expense_rows = conn.execute("""SELECT c.name, COALESCE(SUM(e.amount), 0) FROM expenses e JOIN expense_categories c ON c.category_id = e.category_id
                                       WHERE substr(e.expense_date, 1, 7) = ? AND e.voided_at IS NULL GROUP BY c.name ORDER BY 2 DESC""", (month,)).fetchall()
    expenses = [(a, round(float(b), 2)) for a, b in expense_rows]
    expenses_total = round(sum(v for _, v in expenses), 2)
    net_sales = round(sales_goods - returns, 2)
    net_cogs = round(cogs - returns_cost, 2)
    gross = round(net_sales + services - net_cogs, 2)
    profit = round(gross - expenses_total - shrinkage - lender_interest - bankloan_interest, 2)
    return {"month": month, "sales": net_sales, "sales_gross": sales_goods, "returns": returns, "services": services, "cogs": net_cogs,
            "gross": gross, "expenses": expenses, "expenses_total": expenses_total, "shrinkage": shrinkage,
            "lender_interest": lender_interest, "bankloan_interest": bankloan_interest, "profit": profit}


def xlsx(months: list[dict]) -> bytes:
    """The statement as Excel, months side by side."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Income statement"
    heads = ["", *[m["month"] for m in months]]
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True)
    cats = sorted({name for m in months for name, _ in m["expenses"]})
    def row(label, key=None, values=None, bold=False, sign=1):
        vals = values if values is not None else [sign * m[key] for m in months]
        ws.append([label, *[round(v, 2) for v in vals]])
        if bold:
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
    row("Sales of goods (net of returns)", "sales")
    row("Service income", "services")
    row("Cost of goods sold", "cogs", sign=-1)
    row("Gross margin", "gross", bold=True)
    for cat in cats:
        row(f"  {cat}", values=[-dict(m["expenses"]).get(cat, 0.0) for m in months])
    row("Expenses", "expenses_total", sign=-1, bold=True)
    row("Shrinkage (counts and adjustments)", "shrinkage", sign=-1)
    row("Interest paid to lenders", "lender_interest", sign=-1)
    row("Interest on bank loans", "bankloan_interest", sign=-1)
    row("Profit for the month", "profit", bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 34 if col[0].column == 1 else 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
