"""Sales step S2 (2026-09-29): invoice booklets, the counter sale, the Excel day sheet. Owner's rules: the ERP records the BIR booklet
invoices (never prints); one booklet in use at a time; each number once; spoiled numbers cancelled with a reason; skipped numbers
flagged (never blocking); price agreed per sale, cost = the item's average; below cost allowed and reviewed by the owner; charge only
for customers allowed credit; check payments need bank, number and date; stock may go below zero (flagged). Day sheet: checked line
by line, saved all together or not at all. Defaults (owner: "continue as far as you can"): one payment type per invoice; booklets
registered / closed by the owner."""
import io
import itertools
from datetime import timedelta
from decimal import Decimal

import pytest

pytest.importorskip("django", reason="Django is installed only in the ERP environment (.venv)")

PASSWORD = "sandbox-Test-pass-42"
_SINCE = {"log_id": 0}
_RANGE = itertools.count(20000, 100)


@pytest.fixture
def s2(erp):
    from django.db import connection

    from core.models import Role, User
    from engines.acting_user import acting_as
    from sales import rules
    from sales.models import InvoiceBooklet

    for role in (Role.OWNER, Role.SALES, Role.VIEWER):
        User.objects.filter(email=f"{role}@s2.example.com").delete()
    users = {role: User.objects.create_user(f"{role}@s2.example.com", f"{role.title()} S2", role, PASSWORD, must_change_password=False)
             for role in (Role.OWNER, Role.SALES, Role.VIEWER)}
    with acting_as("owner@s2.example.com"):                 # a clean start: earlier tests' booklets are closed
        for b in InvoiceBooklet.objects.filter(finished_at__isnull=True):
            rules.close_booklet(b, "test set-up")
        first = next(_RANGE)
        booklet = rules.register_booklet(f"T-{first}", str(first), str(first + 49))
    with connection.cursor() as cur:
        cur.execute("SELECT COALESCE(MAX(log_id), 0) FROM audit_logs")
        _SINCE["log_id"] = cur.fetchone()[0]
    return {"users": users, "booklet": booklet, "n": first}


def _audit(action):
    from django.db import connection

    with connection.cursor() as cur:
        cur.execute("SELECT user_email, entity_id, old_values, new_values FROM audit_logs WHERE user_action = %s AND log_id > %s ORDER BY log_id",
                    [action, _SINCE["log_id"]])
        return cur.fetchall()


def _client(user):
    from django.test import Client

    c = Client()
    c.force_login(user)
    return c


def _as(email):
    from engines.acting_user import acting_as

    return acting_as(email)


def _item(brand, qty=10, cost="1000", category="tire", size="185/65R15"):
    from engines.clock import manila_today
    from inventory import stock
    from inventory.models import Item

    from core.models import Party

    with _as("owner@s2.example.com"):
        item = Item.objects.filter(brand=brand, size_model=size).first() or stock.add_item(category, brand, size)
        if qty:
            sup = Party.objects.filter(ap_links__creditor__creditor_type="SUPPLIER").first() or _supplier()
            stock.record_stock_in(sup, manila_today(), "", [(item, qty, cost)])
    item.refresh_from_db()
    return item


def _supplier():
    import sqlite3

    from django.conf import settings

    from core.models import Party
    from core.parties import sync_ap_links

    conn = sqlite3.connect(settings.DATABASES["default"]["NAME"])
    conn.execute("INSERT INTO creditors (creditor_name, creditor_type) VALUES ('S2 Supplier', 'SUPPLIER')")
    conn.commit()
    conn.close()
    sync_ap_links()
    return Party.objects.get(name="S2 Supplier")


def _customer(name, credit_days=None):
    from core.models import Party
    from core.parties import name_key

    p = Party.objects.filter(name=name).first() or Party.objects.create(name=name, name_key=name_key(name), is_customer=True)
    if credit_days is not None:
        p.credit_allowed, p.terms_days = True, credit_days
        p.save()
    return p


def _walk_in():
    from core.models import Party

    return Party.objects.get(is_walk_in=True)


def _sale(n, lines, customer=None, pay="cash", day=None, **kw):
    from engines.clock import manila_today
    from sales import rules

    return rules.SaleInput(number=n, sale_date=day or manila_today(), customer=customer or _walk_in(), payment_type=pay,
                           lines=[rules.SaleLineInput(i, q, p, *r) for i, q, p, *r in lines], **kw)


# ---- booklets -------------------------------------------------------------------------------------------------------------------------

def test_registering_booklets_refuses_overlaps_and_keeps_the_printed_width(s2):
    from sales import rules

    b = s2["booklet"]
    assert b.size == 50 and b.fmt(b.first_number) == str(s2["n"])
    with _as("owner@s2.example.com"):
        with pytest.raises(rules.SaleError, match="overlap"):
            rules.register_booklet("Overlap", str(s2["n"] + 40), str(s2["n"] + 60))
        with pytest.raises(rules.SaleError, match="already registered"):
            rules.register_booklet(b.label, "900001", "900050")
        small = rules.register_booklet(f"Z-{s2['n']}", f"0{s2['n'] + 50}", f"0{s2['n'] + 99}")
    assert small.width == 6 and small.fmt(s2["n"] + 50) == f"0{s2['n'] + 50}"
    assert rules.current_booklet() == b                                   # the oldest unfinished one is in use
    assert _audit("SALES_BOOKLET_REGISTERED")


def test_only_the_booklet_in_use_can_be_used(s2):
    from sales import rules

    item = _item("OneBookletS2")
    with _as("owner@s2.example.com"):
        rules.register_booklet(f"N-{s2['n']}", str(s2["n"] + 50), str(s2["n"] + 99))
        with pytest.raises(rules.SaleError, match="is the booklet in use"):
            rules.record_sale(_sale(s2["n"] + 50, [(item, 1, "1500")]))
        with pytest.raises(rules.SaleError, match="not in any registered booklet"):
            rules.record_sale(_sale(999999, [(item, 1, "1500")]))


def test_cancelling_numbers_needs_a_reason_and_a_finished_booklet_hands_over(s2):
    from sales import rules

    b, n = s2["booklet"], s2["n"]
    with _as("sales@s2.example.com"):
        with pytest.raises(rules.SaleError, match="Say why"):
            rules.cancel_number(n, "")
        rules.cancel_number(n, "spoiled, wrong amount written")
        with pytest.raises(rules.SaleError, match="already cancelled"):
            rules.cancel_number(n, "again")
    assert _audit("SALES_NUMBER_CANCELLED")[0][0] == "sales@s2.example.com"
    with _as("owner@s2.example.com"):
        nxt = rules.register_booklet(f"N2-{n}", str(n + 50), str(n + 99))
        left = rules.close_booklet(b, "lost the last pages")
    b.refresh_from_db()
    assert left == 49 and b.finished_at and rules.current_booklet() == nxt


def test_skipped_numbers_are_flagged_for_the_owner_not_blocked(s2):
    from sales import rules

    item = _item("GapBrandS2")
    n = s2["n"]
    with _as("sales@s2.example.com"):
        rules.record_sale(_sale(n, [(item, 1, "1500")]))
        rules.record_sale(_sale(n + 2, [(item, 1, "1500")]))          # n + 1 skipped: allowed
    assert [(b.label, x) for b, x in rules.gaps(s2["booklet"])] == [(s2["booklet"].label, n + 1)]
    html = _client(s2["users"]["owner"]).get("/").content.decode()
    assert "skipped and not cancelled" in html
    assert "skipped and not cancelled" not in _client(s2["users"]["sales"]).get("/").content.decode()


# ---- a sale ---------------------------------------------------------------------------------------------------------------------------

def test_a_cash_sale_takes_stock_out_at_the_average_cost(s2):
    from sales import rules

    item = _item("CashBrandS2", qty=10, cost="3150")
    with _as("sales@s2.example.com"):
        sale = rules.record_sale(_sale(s2["n"], [(item, 4, "3,800")]))
    item.refresh_from_db()
    line = sale.lines.get()
    assert (sale.total, sale.profit, line.unit_cost) == (Decimal("15200.00"), Decimal("2600.00"), Decimal("3150"))
    assert item.on_hand == 6 and line.stock_move.qty == -4 and line.stock_move.kind == "sale"
    assert sale.customer.is_walk_in and sale.created_by == "sales@s2.example.com"
    assert _audit("SALE_RECORDED")
    with pytest.raises(rules.SaleError, match="already recorded"):
        rules.record_sale(_sale(s2["n"], [(item, 1, "3800")]))


def test_charge_needs_credit_and_gets_a_due_date(s2):
    from engines.clock import manila_today
    from sales import rules

    item = _item("ChargeBrandS2")
    with _as("sales@s2.example.com"):
        with pytest.raises(rules.SaleError, match="not allowed credit"):
            rules.record_sale(_sale(s2["n"], [(item, 1, "1500")], customer=_customer("Cash Only S2"), pay="charge"))
        with pytest.raises(rules.SaleError, match="not allowed credit"):
            rules.record_sale(_sale(s2["n"], [(item, 1, "1500")], pay="charge"))              # walk-in
        sale = rules.record_sale(_sale(s2["n"], [(item, 1, "1500")], customer=_customer("Loyal S2", credit_days=30), pay="charge"))
    assert sale.due_date == manila_today() + timedelta(days=30)


@pytest.mark.parametrize("pay, extra, message", [
    ("check", {}, "bank, check number and check date"),
    ("check", {"check_bank": "BPI", "check_number": "123"}, "bank, check number and check date"),
    ("ewallet", {}, "reference number"),
    ("card", {}, "approval code"),
])
def test_payment_details_are_required(s2, pay, extra, message):
    from sales import rules

    item = _item("PayBrandS2")
    with pytest.raises(rules.SaleError, match=message):
        rules.check_sale(_sale(s2["n"], [(item, 1, "1500")], pay=pay, **extra))


def test_a_post_dated_check_payment_is_kept_with_its_details(s2):
    from engines.clock import manila_today
    from sales import rules

    item = _item("PdcBrandS2")
    with _as("sales@s2.example.com"):
        sale = rules.record_sale(_sale(s2["n"], [(item, 1, "1500")], pay="check", check_bank="BPI", check_number="0045123",
                                       check_date=manila_today() + timedelta(days=15)))
    assert (sale.check_bank, sale.check_number) == ("BPI", "0045123") and sale.check_date > manila_today()


def test_below_cost_is_allowed_and_goes_to_the_owners_review(s2):
    from sales import rules
    from sales.models import SaleLine

    item = _item("BelowBrandS2", cost="3050")
    with _as("sales@s2.example.com"):
        chk = rules.check_sale(_sale(s2["n"], [(item, 1, "2900", "keep customer")]))
        assert any("below cost" in w for w in chk["warnings"])
        sale = rules.record_sale(_sale(s2["n"], [(item, 1, "2900", "keep customer")]))
    line = sale.lines.get()
    assert line.below_cost and line.below_cost_reason == "keep customer" and line.profit == Decimal("-150.00")
    owner, sales = _client(s2["users"]["owner"]), _client(s2["users"]["sales"])
    assert line in list(owner.get("/sales/below-cost/").context["lines"])
    assert sales.get("/sales/below-cost/").status_code == 403
    assert sales.post(f"/sales/below-cost/{line.pk}/reviewed/").status_code == 403
    owner.post(f"/sales/below-cost/{line.pk}/reviewed/")
    line.refresh_from_db()
    assert line.reviewed_by == "owner@s2.example.com" and _audit("SALE_BELOW_COST_REVIEWED")
    assert line not in list(owner.get("/sales/below-cost/").context["lines"])


def test_selling_more_than_on_hand_is_allowed_and_flagged(s2):
    from sales import rules

    item = _item("ShortBrandS2", qty=1)
    with _as("sales@s2.example.com"):
        assert any("below zero" in w for w in rules.check_sale(_sale(s2["n"], [(item, 3, "1500")]))["warnings"])
        rules.record_sale(_sale(s2["n"], [(item, 3, "1500")]))
    item.refresh_from_db()
    assert item.on_hand == -2 and "less than zero on hand" in _client(s2["users"]["owner"]).get("/").content.decode()


@pytest.mark.parametrize("case, message", [
    ("future", "future"), ("no_lines", "at least one item"), ("zero_qty", "at least 1"), ("half_qty", "whole pieces"),
    ("negative_price", "can't be negative"), ("twice", "on two lines"), ("inactive_customer", "not an active customer"),
])
def test_sale_mistakes_are_refused_and_nothing_is_saved(s2, case, message):
    from engines.clock import manila_today
    from sales import rules
    from sales.models import Sale

    item = _item(f"MistakeS2 {case}")
    lines = {"no_lines": [], "zero_qty": [(item, 0, "10")], "half_qty": [(item, "1.5", "10")], "negative_price": [(item, 1, "-1")],
             "twice": [(item, 1, "10"), (item, 1, "10")]}.get(case, [(item, 1, "10")])
    customer = None
    if case == "inactive_customer":
        customer = _customer("Closed Shop S2")
        customer.is_active = False
        customer.save()
    day = manila_today() + timedelta(days=1) if case == "future" else None
    with pytest.raises(rules.SaleError, match=message):
        rules.record_sale(_sale(s2["n"], lines, customer=customer, day=day))
    assert not Sale.objects.filter(number=s2["n"]).exists()
    item.refresh_from_db()
    assert item.on_hand == 10


def test_the_quick_sale_page_checks_then_saves(s2):
    from sales.models import Sale

    item = _item("PageBrandS2", qty=5, cost="2000")
    sales = _client(s2["users"]["sales"])
    page = sales.get("/sales/new/")
    assert page.context["f"]["number"] == str(s2["n"])                   # the next unused number is suggested
    data = {"number": str(s2["n"]), "sale_date": page.context["today"], "customer": str(_walk_in().pk), "payment_type": "cash",
            "item_0": f"{item.code} · {item.description}", "qty_0": "2", "price_0": "2,500", "action": "check"}
    resp = sales.post("/sales/new/", data)
    assert resp.context["chk"]["profit"] == Decimal("1000.00") and not Sale.objects.filter(number=s2["n"]).exists()
    data["action"] = "save"
    resp = sales.post("/sales/new/", data)
    sale = Sale.objects.get(number=s2["n"])
    assert resp["Location"] == f"/sales/{sale.pk}/" and sale.total == Decimal("5000.00")
    day = sales.get("/sales/", {"from": page.context["today"]})
    assert sale in list(day.context["sales"]) and day.context["pieces"] >= 2


# ---- the Excel day sheet --------------------------------------------------------------------------------------------------------------

def _sheet(rows, cancelled=(), day=None):
    from openpyxl import load_workbook

    from engines.clock import manila_today
    from sales import daysheet

    wb = load_workbook(io.BytesIO(daysheet.build_template()))
    ws = wb["Day"]
    ws["B1"] = day or manila_today()
    for r, row in enumerate(rows, start=daysheet.FIRST):
        for col, v in zip("ABCGHIJKLMN", row):
            ws[f"{col}{r}"] = v
    for r, (no, reason) in enumerate(cancelled, start=2):
        wb["Cancelled"][f"A{r}"], wb["Cancelled"][f"B{r}"] = no, reason
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    out.name = "day.xlsx"
    return out


def test_the_day_sheet_template_carries_todays_lists(s2):
    from openpyxl import load_workbook

    from sales import daysheet

    item = _item("TemplateBrandS2")
    wb = load_workbook(io.BytesIO(daysheet.build_template()))
    assert wb.sheetnames == ["How to fill", "Day", "Lists", "Cancelled"] and wb["Lists"].sheet_state == "hidden"
    assert [c.value for c in wb["Day"][4]][:8] == ["Qty", "Item", "Price each", "Cost each", "Amount", "Profit", "Customer", "Invoice no."]
    lists = wb["Lists"]
    assert daysheet.item_choice(item) in [c.value for c in lists["A"]]
    assert str(s2["n"]) == lists["H2"].value                               # next unused number first
    assert wb["Day"]["F5"].value.startswith("=IF(")                         # profit fills in by formula
    assert "INDEX(" in wb["Day"]["D5"].value and "XLOOKUP" not in wb["Day"]["D5"].value   # XLOOKUP stayed blank in real Excel


def test_a_clean_day_sheet_saves_every_sale_together(s2):
    from sales.models import CancelledNumber, Sale

    a, b = _item("DayA S2", qty=10, cost="3150"), _item("DayB S2", qty=5, cost="4050", category="battery", size="NS60")
    loyal = _customer("Juan Trucking S2", credit_days=30)
    n = s2["n"]
    f = _sheet([(2, a.code + " · x", 3800, "", str(n), "Cash"),
                (4, a.code + " · x", 3800, loyal.name, str(n + 1), "Charge (credit)"),
                (1, b.code + " · x", 5200, loyal.name, str(n + 1), "Charge (credit)")],
               cancelled=[(str(n + 2), "spoiled")])
    sales = _client(s2["users"]["sales"])
    resp = sales.post("/sales/day-sheet/", {"file": f})
    parsed = resp.context["parsed"]
    assert parsed["errors"] == [] and parsed["count"] == 2 and parsed["total"] == Decimal("28000.00")
    assert not Sale.objects.filter(number__in=[n, n + 1]).exists()          # nothing saved by the check
    resp = sales.post("/sales/day-sheet/", {"action": "save"})
    two = Sale.objects.get(number=n + 1)
    assert two.lines.count() == 2 and two.payment_type == "charge" and two.upload is not None
    assert CancelledNumber.objects.filter(number=n + 2, reason="spoiled").exists()
    a.refresh_from_db()
    assert a.on_hand == 4 and _audit("SALES_DAY_UPLOADED")


def test_day_sheet_mistakes_are_listed_and_nothing_is_saved(s2):
    from sales.models import Sale

    a = _item("DayErr S2")
    n = s2["n"]
    f = _sheet([(1, a.code + " · x", 100, "", "", "Cash"),                              # no invoice number
                (1, "ZZ-9999 · nothing", 100, "", str(n), "Cash"),                       # unknown item
                (1, a.code + " · x", 100, "Nobody S2", str(n + 1), "Cash"),              # unknown customer
                (1, a.code + " · x", 100, "", str(n + 2), "Check"),                      # check without details
                (1, a.code + " · x", 100, "", str(n + 3), "Cash"),
                (1, _item("DayErr2 S2").code + " · x", 100, "", str(n + 3), "Card")])    # two payments on one invoice
    parsed = _client(s2["users"]["sales"]).post("/sales/day-sheet/", {"file": f}).context["parsed"]
    text = " ".join(parsed["errors"])
    assert "pick the invoice number" in text and "pick the item" in text and "not on the customer list" in text
    assert "bank, check number and check date" in text and "different payment" in text
    assert not Sale.objects.filter(number__gte=n, number__lte=n + 3).exists()


def test_the_day_is_all_or_nothing_even_if_something_changed_after_the_check(s2):
    from sales import rules
    from sales.models import Sale

    a = _item("Race S2")
    n = s2["n"]
    sales = _client(s2["users"]["sales"])
    sales.post("/sales/day-sheet/", {"file": _sheet([(1, a.code + " · x", 100, "", str(n), "Cash"),
                                                     (1, a.code + " · x", 100, "", str(n + 1), "Cash")])})
    with _as("owner@s2.example.com"):                                       # someone uses n + 1 before "Save the day"
        rules.record_sale(_sale(n + 1, [(a, 1, "100")]))
    resp = sales.post("/sales/day-sheet/", {"action": "save"}, follow=True)
    assert "Nothing was saved" in resp.content.decode()
    assert not Sale.objects.filter(number=n).exists()


# ---- who sees what --------------------------------------------------------------------------------------------------------------------

def test_every_sales_page_declares_roles_and_viewers_are_refused(s2):
    from sales import urls

    viewer = _client(s2["users"]["viewer"])
    for pattern in urls.urlpatterns:
        assert getattr(pattern.callback, "allowed_roles", None), pattern.name
    for path in ("/sales/", "/sales/new/", "/sales/day-sheet/", "/sales/day-sheet/template/", "/sales/booklets/", "/sales/below-cost/"):
        assert viewer.get(path).status_code == 403, path


def test_only_the_owner_registers_or_closes_booklets(s2):
    sales = _client(s2["users"]["sales"])
    assert sales.post("/sales/booklets/register/", {"label": "X", "first": "1", "last": "50"}).status_code == 403
    assert sales.post(f"/sales/booklets/{s2['booklet'].pk}/close/", {"note": "x"}).status_code == 403
    assert sales.get("/sales/booklets/").status_code == 200


def test_only_the_closest_menu_item_is_highlighted(s2):
    sales = _client(s2["users"]["sales"])
    item = _item("MenuBrandS2", qty=0)
    for path, label in (("/sales/booklets/", "Invoice booklets"), ("/sales/", "Sales"), ("/sales/new/", "New sale"),
                        (f"/items/{item.pk}/", "Items and stock"), ("/stock-in/new/", "Stock in")):
        nav = sales.get(path).context["nav"]
        assert [i["label"] for i in nav if i["on"]] == [label], path
