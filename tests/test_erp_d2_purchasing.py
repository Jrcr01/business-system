"""Django-only move, step D2 (2026-10-02): the Purchasing screens (price lists, purchase orders, receiving), by role, on the shared rules.
Purchasing makes and sends orders (their own), prices come from the Excel template with the supplier's discount; Payables receives the invoice into an AP batch and the
goods into stock at the price actually charged; a role that may not do a thing is refused, even through the address bar."""
import io
import re
import sqlite3
from datetime import date

import pytest

pytest.importorskip("django", reason="Django is installed only in the ERP environment (.venv)")

from openpyxl import load_workbook  # noqa: E402

PASSWORD = "sandbox-Test-pass-42"


def _db():
    from django.conf import settings

    return sqlite3.connect(settings.DATABASES["default"]["NAME"])


@pytest.fixture
def d2(erp):
    from core.models import Role, User
    from engines import discount_rules as dr

    users = {}
    for role in (Role.OWNER, Role.PURCHASING, Role.PAYABLES, Role.VIEWER):
        User.objects.filter(email=f"{role}@d2.example.com").delete()
        users[role] = User.objects.create_user(f"{role}@d2.example.com", f"{role.title()} D2", role, PASSWORD, must_change_password=False)
    c = _db()
    cid = c.execute("INSERT INTO creditors (creditor_name, creditor_type, terms) VALUES ('D2 Supplier', 'SUPPLIER', '30')").lastrowid
    c.commit()
    c.close()
    users["cid"] = cid
    from engines.acting_user import acting_as
    with acting_as("owner@d2.example.com", "owner"):
        dr.set_discount(cid, "D2Westlake", "5,5", date(2026, 9, 1))
    return users


def _client(user):
    from django.test import Client

    c = Client()
    c.force_login(user)
    return c


def _template(cid):
    """A filled price template (one tire at 7,600) as an upload."""
    from engines import price_template_service as pt

    wb = load_workbook(io.BytesIO(pt.build_template()))
    ws = wb["Tires"]
    head = [str(c.value).replace(" *", "") for c in ws[1]]
    for k, v in {"Brand": "D2Westlake", "Size": "285/50R20", "Terrain": "A/T", "Pattern": "SL369", "Supplier item code": "WL-1", "Unit": "pc", "List price": 7600,
                 "Effective date": date(2026, 9, 15)}.items():
        ws.cell(row=2, column=head.index(k) + 1, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "prices.xlsx"
    return buf


def test_purchasing_loads_prices_from_the_template_and_sees_the_net_after_the_discount(d2):
    from core.models import Role

    c, cid = _client(d2[Role.PURCHASING]), d2["cid"]
    page = c.get(f"/price-lists/?supplier={cid}&tab=import")
    assert page.status_code == 200 and b"Blank template" in page.content
    assert c.get("/price-lists/template/").status_code == 200
    r = c.post("/price-lists/", {"tab": "import", "supplier": cid, "step": "preview", "start": "2026-09-15", "file": _template(cid)})
    assert r.status_code == 200 and b"1</b> new" in r.content
    c.post("/price-lists/", {"tab": "import", "supplier": cid, "step": "confirm", "dup": "skip"})
    page = c.get(f"/price-lists/?supplier={cid}&tab=current&on=2026-10-01")
    html = page.content.decode()
    assert "D2Westlake 285/50R20 (A/T) SL369" in html and "7,600.00" in html and "6,859.00" in html and "less 5% + 5%" in html


def test_a_viewer_sees_prices_but_has_no_import_or_discount_tabs_and_cannot_change_them(d2):
    from core.models import Role

    c, cid = _client(d2[Role.VIEWER]), d2["cid"]
    html = c.get(f"/price-lists/?supplier={cid}").content.decode()
    assert "Import Excel" not in html and "Discounts" not in html
    c.post("/price-lists/", {"tab": "discounts", "supplier": cid, "brand": "X", "chain": "50", "start": "2026-10-01"})
    c2 = _db()
    assert c2.execute("SELECT COUNT(*) FROM supplier_discounts WHERE brand = 'X'").fetchone()[0] == 0


def test_the_whole_cycle_order_by_purchasing_receive_by_payables_stock_and_batch(d2):
    from core.models import Role
    from inventory.models import Item

    cid = d2["cid"]
    buyer, clerk, owner = _client(d2[Role.PURCHASING]), _client(d2[Role.PAYABLES]), _client(d2[Role.OWNER])
    owner.post("/price-lists/", {"tab": "import", "supplier": cid, "step": "preview", "start": "2026-09-15", "file": _template(cid)})
    owner.post("/price-lists/", {"tab": "import", "supplier": cid, "step": "confirm", "dup": "skip"})
    item = Item.objects.get(brand="D2Westlake")

    # Purchasing: the form shows the payables panel, then saves and orders
    form = buyer.get(f"/purchase-orders/new/?supplier={cid}")
    assert form.status_code == 200 and b"Payables when this order falls due" in form.content
    post = {"action": "save_order", "supplier": cid, "order_date": "2026-10-01", "delivery": "2026-10-08", "notes": "",
            "item_0": f"{item.code} · {item.name}", "qty_0": "4", "price_0": "", "disc_0": ""}
    r = buyer.post("/purchase-orders/new/", post)
    assert r.status_code == 302, r.content.decode()[:500]
    detail = buyer.get(r["Location"])
    html = detail.content.decode()
    assert "6,859.00" in html and "27,436.00" in html and "Ordered" in html and "Download PO (PDF)" in html
    po_id = int(re.search(r"/purchase-orders/(\d+)/", r["Location"]).group(1))
    assert buyer.get(f"/purchase-orders/{po_id}/pdf/")["Content-Type"] == "application/pdf"
    assert "Receive / match invoice" not in html                                           # receiving is not Purchasing's
    assert buyer.get("/receive/").status_code == 403

    # Payables: receive at the real price, into a new batch and into stock
    rcv = clerk.get(f"/receive/?supplier={cid}&po={po_id}")
    assert rcv.status_code == 200 and b"Still to come" in rcv.content
    line = _db().execute("SELECT line_id FROM purchase_order_lines WHERE po_id = ?", (po_id,)).fetchone()[0]
    r = clerk.post("/receive/", {"step": "receive", "supplier": cid, "invoice_no": "SI-D2-1", "invoice_date": "2026-10-09", "batch": "0",
                                 f"pick_{line}": "on", f"qty_{line}": "4", f"price_{line}": "6900", "update_prices": "on"})
    assert r.status_code == 302
    c2 = _db()
    assert c2.execute("SELECT gross_amount, due_date FROM invoices WHERE doc_ref = 'SI-D2-1'").fetchone() == (27600.0, "2026-11-09")                 # invoice date + 30 days = Sunday Nov 8: due on the next bank day
    assert c2.execute("SELECT status FROM purchase_orders WHERE po_id = ?", (po_id,)).fetchone()[0] == "RECEIVED"
    item.refresh_from_db()
    assert item.on_hand == 4 and float(item.avg_cost) == 6900.0                            # the goods went into stock at the price charged
    # the list price was worked back from the net 6,900: 6,900 / 0.9025
    assert c2.execute("SELECT unit_price FROM supplier_price_list WHERE item_id = ? ORDER BY price_id DESC LIMIT 1", (item.pk,)).fetchone()[0] == 7645.43


def test_a_purchasing_user_cannot_change_someone_elses_order_or_cancel_it(d2):
    from core.models import Role

    cid = d2["cid"]
    owner, buyer = _client(d2[Role.OWNER]), _client(d2[Role.PURCHASING])
    owner.post("/purchase-orders/new/", {"action": "save", "supplier": cid, "order_date": "2026-10-01", "delivery": "2026-10-08",
                                         "item_0": "Mystery part", "qty_0": "1", "price_0": "100", "disc_0": "0"})
    po_id = _db().execute("SELECT MAX(po_id) FROM purchase_orders").fetchone()[0]
    assert buyer.get(f"/purchase-orders/{po_id}/edit/").status_code == 302                  # sent back to the list with a message
    buyer.post(f"/purchase-orders/{po_id}/", {"action": "cancel"})
    assert _db().execute("SELECT status FROM purchase_orders WHERE po_id = ?", (po_id,)).fetchone()[0] == "DRAFT"
    owner.post(f"/purchase-orders/{po_id}/", {"action": "cancel", "reason": "test"})
    assert _db().execute("SELECT status FROM purchase_orders WHERE po_id = ?", (po_id,)).fetchone()[0] == "CANCELLED"


def test_picking_a_lighter_week_moves_the_delivery_date(d2):
    from core.models import Role

    cid = d2["cid"]
    buyer = _client(d2[Role.PURCHASING])
    base = {"supplier": cid, "order_date": "2026-10-01", "delivery": "2026-10-08", "notes": ""}
    page = buyer.post("/purchase-orders/new/", {**base, "action": "refresh"}).content.decode()
    weeks = re.findall(r'value="week:(\d{4}-\d{2}-\d{2})"\s*(?:style="[^"]*")?\s*>', page) or re.findall(r'value="week:(\d{4}-\d{2}-\d{2})"', page)
    assert weeks, "the week buttons are shown"
    target = weeks[-1]
    moved = buyer.post("/purchase-orders/new/", {**base, "action": f"week:{target}"}).content.decode()
    new_delivery = re.search(r'name="delivery" type="date" value="(\d{4}-\d{2}-\d{2})"', moved).group(1)
    assert new_delivery != "2026-10-08"
