"""Stage D2a of the Django rebuild (Price lists, docs/mockups/d2_buying.html screens 1-3): one supplier's prices and their history, add a price, the Excel template
(check the file first, then import), discounts, and the catalog with suppliers side by side. Payables and the Viewer may look; only Purchasing, the owners and the
System admin change prices."""
import io

from tests.conftest import fetchone, make_creditor
from tests.test_web_d0 import _admin, _sign_in_new


def _db():
    import engines.db_engine as d
    return d.DB_PATH


def _price(web, cid, item, price, start="2026-09-01", **extra):
    return web.post("/price-lists/", {"s": cid, "action": "add_price", "item": item, "unit": "pc", "price": price, "start": start, "code": "", **extra})


def test_add_a_price_and_see_it_with_the_change(web, db):
    _admin(web)
    cid = make_creditor(db, "Goodyear Philippines")
    assert _price(web, cid, "Eagle F1 205/55 R16", "7,300")["Location"].startswith("/price-lists/?s=")
    _price(web, cid, "Eagle F1 205/55 R16", "7600", start="2026-10-01")
    page = web.get(f"/price-lists/?s={cid}&tab=prices&on=2026-10-03").content.decode()
    assert "Eagle F1 205/55 R16" in page and "7,600.00" in page and "7,300.00" in page and "+4.1%" in page
    older = web.get(f"/price-lists/?s={cid}&tab=prices&on=2026-09-15").content.decode()
    assert "7,300.00" in older and "7,600.00" not in older                       # prices in force on an earlier day


def test_history_and_switching_a_wrong_price_off(web, db):
    _admin(web)
    cid = make_creditor(db, "Goodyear Philippines")
    _price(web, cid, "Assurance 185/65 R15", "4950")
    pid = fetchone(_db(), "SELECT price_id FROM supplier_price_list WHERE creditor_id = ?", (cid,))[0]
    page = web.get(f"/price-lists/?s={cid}&tab=prices&item=Assurance%20185/65%20R15").content.decode()
    assert "History of Assurance 185/65 R15" in page and "Switch off" in page
    web.post("/price-lists/", {"s": cid, "action": "price_on_off", "price_id": pid, "on": "0", "item": "Assurance 185/65 R15"})
    assert fetchone(_db(), "SELECT is_active FROM supplier_price_list WHERE price_id = ?", (pid,))[0] == 0
    assert "4,950.00" not in web.get(f"/price-lists/?s={cid}&tab=prices").content.decode().split("History of")[0]


def test_discounts_are_set_and_listed(web, db):
    _admin(web)
    cid = make_creditor(db, "Tire Depot")
    r = web.post("/price-lists/", {"s": cid, "action": "set_discount", "brand": "", "chain": "5, 5", "start": "2026-01-01"})
    assert r.status_code == 302
    page = web.get(f"/price-lists/?s={cid}&tab=discounts").content.decode()
    assert "All brands" in page and "less 5%, less 5%" in page


def test_the_template_downloads_and_an_import_is_checked_before_it_is_saved(web, db):
    from openpyxl import load_workbook

    from engines import price_template_service as pt
    _admin(web)
    cid = make_creditor(db, "Motolite")
    blank = web.get("/price-lists/template/")
    assert blank.status_code == 200 and blank["Content-Type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(blank.content))
    ws = wb["Battery"]
    keys = [k for k, _, _ in pt._headers("Battery")]
    row = {"brand": "Motolite", "model": "NS60", "volts": "12", "ah": "45", "terminal": "L", "list_price": 5200}
    for j, k in enumerate(keys, start=1):
        if k in row:
            ws.cell(row=2, column=j, value=row[k])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "motolite.xlsx"
    r = web.post("/price-lists/", {"s": cid, "action": "check_file", "start": "2026-10-01", "file": buf})
    assert r.status_code == 302 and "file=" in r["Location"]
    check = web.get(r["Location"]).content.decode()
    assert "before anything is saved" in check and "Import 1 row(s)" in check
    assert fetchone(_db(), "SELECT COUNT(*) FROM supplier_price_list WHERE creditor_id = ?", (cid,))[0] == 0     # nothing saved by the check
    token = r["Location"].split("file=")[1].split("&")[0]
    done = web.post("/price-lists/", {"s": cid, "action": "import_file", "file": token, "start": "2026-10-01", "dup": "skip"})
    assert done.status_code == 302
    assert fetchone(_db(), "SELECT COUNT(*), MAX(unit_price) FROM supplier_price_list WHERE creditor_id = ?", (cid,)) == (1, 5200)
    assert fetchone(_db(), "SELECT COUNT(*) FROM items WHERE brand = 'Motolite'")[0] == 1                      # the product got its SKU


def test_a_wrong_file_is_refused(web, db):
    _admin(web)
    cid = make_creditor(db, "Motolite")
    bad = io.BytesIO(b"not excel")
    bad.name = "prices.csv"
    r = web.post("/price-lists/", {"s": cid, "action": "check_file", "file": bad})
    assert "file=" not in r["Location"]


def test_catalog_add_discontinue_and_side_by_side(web, db):
    _admin(web)
    a, b = make_creditor(db, "Supplier A"), make_creditor(db, "Supplier B")
    r = web.post("/price-lists/", {"s": a, "action": "add_product", "cat": "Battery", "brand": "Motolite", "model": "NS60", "ah": "45", "terminal": "L"})
    assert r.status_code == 302
    name, pid = fetchone(_db(), "SELECT name, id FROM items WHERE brand = 'Motolite'")
    _price(web, a, name, "5200")
    _price(web, b, name, "5000")
    side = web.get(f"/price-lists/?s={a}&tab=catalog&cat=Battery").content.decode()
    assert "Supplier A" in side and "Supplier B" in side and "cheapest" in side
    web.post("/price-lists/", {"s": a, "action": "product_on_off", "product_id": pid, "cat": "Battery", "on": "0"})
    assert fetchone(_db(), "SELECT is_active FROM items WHERE id = ?", (pid,))[0] == 0


def test_payables_may_look_but_not_change(web, db):
    _admin(web)
    cid = make_creditor(db, "Goodyear Philippines")
    _price(web, cid, "Eagle F1", "7600")
    payables = _sign_in_new("ap@slicks.test", "payables")
    page = payables.get(f"/price-lists/?s={cid}").content.decode()
    assert "Eagle F1" in page and "Add a price" not in page and "Import Excel" not in page
    assert payables.post("/price-lists/", {"s": cid, "action": "add_price", "item": "X", "unit": "pc", "price": "1", "start": "2026-10-01"}).status_code == 403
    assert fetchone(_db(), "SELECT COUNT(*) FROM supplier_price_list WHERE item_name = 'X'")[0] == 0
    mother = _sign_in_new("mother@slicks.test", "releaser")
    assert mother.get("/price-lists/").status_code == 403


def test_no_supplier_yet(web):
    _admin(web)
    assert b"No supplier yet" in web.get("/price-lists/").content
