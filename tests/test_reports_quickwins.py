"""Reports review quick wins (2026-09-28, plan phase 3): friendly column names and CSV / Excel exports, the Schedule of payments filters
(date range, payees) shared by screen and PDF, the rundown read once and narrowed in memory, the shared per-supplier pivot, and PDFs
that survive "&" in names and use the caller's column widths."""
import io
import sqlite3
from datetime import date, timedelta

import openpyxl
import pandas as pd
import pypdfium2 as pdfium

import app
from conftest import make_creditor, make_statement
from engines.account_history_pdf_service import generate_rundown_pdf
from engines.export_service import csv_bytes, excel_bytes, friendly, label
from engines.payment_schedule_report_service import generate_payment_schedule_pdf
from engines.reports_service import filter_rundown, filter_schedule, get_payables_rundown, get_payment_schedule_all_tracks
from views.reports import _status_pivot

SAMPLE = pd.DataFrame({"check_number": ["3992000001", "3992000002"], "check_date": ["2026-10-01", "2026-10-15"],
                       "creditor_name": ["S&P Tires", "ACME"], "amount": [1234.5, 99.999], "mystery_col": ["a", "b"]})


def test_friendly_names_and_csv():
    assert label("check_date") == "Check date" and label("creditor_name") == "Payee" and label("mystery_col") == "Mystery col"
    assert label("amount", {"amount": "Net amount"}) == "Net amount"
    assert list(friendly(SAMPLE).columns) == ["Check #", "Check date", "Payee", "Amount", "Mystery col"]
    text = csv_bytes(SAMPLE).decode("utf-8-sig").splitlines()
    assert text[0] == "Check #,Check date,Payee,Amount,Mystery col"
    assert text[1] == "3992000001,2026-10-01,S&P Tires,1234.50,a" and text[2].endswith(",100.00,b")


def test_excel_has_one_sheet_per_table_with_formats():
    data = excel_bytes({"Checks": SAMPLE, "Per supplier": pd.DataFrame({"supplier": ["ACME"], "Total": [5.0]})})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Checks", "Per supplier"]
    ws = wb["Checks"]
    assert [c.value for c in ws[1]] == ["Check #", "Check date", "Payee", "Amount", "Mystery col"]
    assert ws["B2"].number_format == "mmm dd, yyyy" and ws["B2"].value.date() == date(2026, 10, 1)
    assert ws["D2"].number_format == "#,##0.00" and ws["D2"].value == 1234.5
    assert ws.freeze_panes == "A2"


def _schedule_data(db):
    a, b = make_creditor(db, name="S&P Tires"), make_creditor(db, name="ACME")
    today = date.today()
    for cid, days, amt in ((a, 3, 100.0), (a, 40, 200.0), (b, 5, 300.0)):
        sid = make_statement(db, cid)
        with sqlite3.connect(db) as conn:                   # make_statement always uses 'TEST-REF'
            conn.execute("UPDATE order_statements SET statement_ref = ?, status = 'Issued' WHERE statement_id = ?", (f"OS-{cid}-{days}", sid))
        app.add_invoice(cid, f"INV-{cid}-{days}", today, today + timedelta(days=days), amt, [], sid)
        app.issue_check(cid, amt, today + timedelta(days=days), statement_id=sid)
    return today


def test_schedule_filters_match_on_screen_and_pdf(db):
    today = _schedule_data(db)
    full = get_payment_schedule_all_tracks()
    assert len(full) == 3
    assert filter_schedule(full, payees=["S&P Tires"])["amount"].sum() == 300.0
    near = filter_schedule(full, date_from=today, date_to=today + timedelta(days=10))
    assert sorted(near["amount"]) == [100.0, 300.0]
    assert filter_schedule(full, ["Supplier"], today, today + timedelta(days=10), ["ACME"])["amount"].tolist() == [300.0]
    assert filter_schedule(full, ["Bank Loan"]).empty
    pdf = generate_payment_schedule_pdf(["Supplier"], None, today, today + timedelta(days=10), ["S&P Tires"])
    doc = pdfium.PdfDocument(pdf)
    text = " ".join(" ".join(doc[i].get_textpage().get_text_range().split()) for i in range(len(doc)))
    assert "S&P TIRES" in text and "100.00" in text and "300.00" not in text


def test_rundown_filtered_in_memory_matches_the_query(db):
    today = _schedule_data(db)
    everything = get_payables_rundown(None)
    lo, hi = today, today + timedelta(days=10)
    pd.testing.assert_frame_equal(filter_rundown(everything, lo, hi).reset_index(drop=True),
                                  get_payables_rundown(None, lo, hi).reset_index(drop=True))
    assert len(filter_rundown(everything, lo, hi)) == 2 and filter_rundown(everything) is everything


def test_status_pivot_fills_missing_statuses():
    df = pd.DataFrame({"supplier": ["A", "A", "B"], "status": ["Paid", "Pending", "Paid"], "net_total": [10.0, 5.0, 7.0]})
    piv = _status_pivot(df, "net_total")
    assert list(piv.columns) == ["supplier", "Paid", "Partially Paid", "Pending", "Total"]
    assert piv.set_index("supplier").loc["A"].tolist() == [10.0, 0.0, 5.0, 15.0]


def test_rundown_pdf_with_ampersands_and_long_text():
    rows = [[f"S&P Tires and Rubber Supply Corporation {i} <branch>", f"BATCH-{i}", "Pending", "2026-10-01", "1,234.00", "a very long note " * 6]
            for i in range(30)]
    pdf = generate_rundown_pdf("Batch & rundown", "S&P Tires", ["Supplier", "Batch", "Status", "Date", "Amount", "Note"], rows,
                               [0.35, 0.1, 0.08, 0.08, 0.09, 0.30], [4], [("Pending", 30, 37020.0)], "Checks dated 2026 & later")
    doc = pdfium.PdfDocument(pdf)
    text = " ".join(" ".join(doc[i].get_textpage().get_text_range().split()) for i in range(len(doc)))
    assert "BATCH & RUNDOWN" in text and "S&P Tires" in text and "<branch>" in text
