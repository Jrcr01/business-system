""""My reports" builder (2026-09-28, plan phase 4): only listed data sets / fields / operations, relative dates that roll forward, every
filter type, grouping with subtotals and a grand total, sort and Top N, exports (CSV / Excel / PDF), and saved reports (save, re-run,
duplicate, delete, pin to the Dashboard, include in the overall report)."""
import io
from datetime import date

import openpyxl
import pandas as pd
import pypdfium2 as pdfium
import pytest

from engines import report_builder as rb
from engines.export_service import csv_bytes, excel_bytes

TODAY = date(2026, 9, 28)            # a Monday

CHECKS = pd.DataFrame([
    # check_number, check_date, payee, track, reference, amount, status, printed, bank_cleared_date, days_to_clear
    ("3992000001", "2026-09-01", "ACME", "Supplier", "OS-1", 1000.0, "Active", "Yes", "2026-09-03", 2),
    ("3992000002", "2026-09-15", "ACME", "Supplier", "OS-1", 3000.0, "Active", "Yes", None, None),
    ("3992000003", "2026-09-20", "S&P Tires", "Supplier", "OS-2", 500.0, "Active", "No", None, None),
    ("3992000004", "2026-08-10", "S&P Tires", "Supplier", "OS-3", 700.0, "Cancelled", "Yes", None, None),
    ("3992000005", "2026-10-05", "Lender Co", "Private Lender", "LN-1", 2000.0, "Active", "No", None, None),
], columns=list(rb.DATASETS["checks"].fields))


def spec(**kw):
    s = rb.blank_spec("checks")
    s.update(kw)
    return s


def test_presets_are_worked_out_on_the_day():
    r = rb.preset_range
    assert r("This week", TODAY) == (date(2026, 9, 28), date(2026, 10, 4))
    assert r("Last week", TODAY) == (date(2026, 9, 21), date(2026, 9, 27))
    assert r("This month", TODAY) == (date(2026, 9, 1), date(2026, 9, 30))
    assert r("Last month", TODAY) == (date(2026, 8, 1), date(2026, 8, 31))
    assert r("This quarter", TODAY) == (date(2026, 7, 1), date(2026, 9, 30))
    assert r("Last quarter", TODAY) == (date(2026, 4, 1), date(2026, 6, 30))
    assert r("Last year", TODAY) == (date(2025, 1, 1), date(2025, 12, 31))
    assert r("Last N days", TODAY, 7) == (date(2026, 9, 22), TODAY)
    assert r("Next N days", TODAY, 30) == (TODAY, date(2026, 10, 28))
    assert r("Custom range", TODAY, None, "2026-01-01", None) == (date(2026, 1, 1), None)
    assert r("All dates", TODAY) == (None, None)
    s = spec(preset="This month")                                     # the same saved report, a month later, shows October
    assert list(rb.run(s, today=TODAY, data=CHECKS).table["Check #"]) == ["3992000001", "3992000002", "3992000003"]
    assert list(rb.run(s, today=date(2026, 10, 20), data=CHECKS).table["Check #"]) == ["3992000005"]


def test_only_listed_fields_and_operations_are_allowed():
    assert rb.validate(spec()) == []
    assert rb.validate({"dataset": "users; DROP TABLE checks_issued"})
    assert rb.validate(spec(columns=["check_number", "card_number_full"]))
    assert rb.validate(spec(filters=[{"field": "amount", "op": "contains", "value": "1"}]))
    assert rb.validate(spec(group_by=[{"field": "payee", "part": "month"}]))           # only dates have months
    assert rb.validate(spec(aggregates=[{"fn": "sum", "field": "payee"}]))
    assert rb.validate(spec(group_by=[{"field": "payee"}, {"field": "track"}, {"field": "status"}]))
    with pytest.raises(ValueError):
        rb.run(spec(columns=["nope"]), data=CHECKS)
    for ds in rb.DATASETS.values():                                    # full card numbers are never offered
        assert "card_number_full" not in (ds.sql or "") and "card_number_full" not in ds.fields


def test_every_filter_type():
    run = lambda *f: list(rb.run(spec(filters=list(f)), data=CHECKS).table["Check #"])  # noqa: E731
    assert run({"field": "payee", "op": "contains", "value": "s&p"}) == ["3992000004", "3992000003"]
    assert run({"field": "payee", "op": "is one of", "value": ["ACME"]}) == ["3992000001", "3992000002"]
    assert len(run({"field": "status", "op": "is not one of", "value": ["Cancelled"]})) == 4
    assert run({"field": "amount", "op": "between", "value": [900, 2000]}) == ["3992000001", "3992000005"]
    assert run({"field": "amount", "op": "between", "value": [None, 600]}) == ["3992000003"]
    assert run({"field": "bank_cleared_date", "op": "between", "value": ["2026-09-01", None]}) == ["3992000001"]
    r = rb.run(spec(filters=[{"field": "payee", "op": "is one of", "value": ["ACME"]}, {"field": "amount", "op": "between", "value": [2000, None]}]),
               data=CHECKS)
    assert list(r.table["Check #"]) == ["3992000002"] and r.headline == ("Total amount", 3000.0, "money")
    assert "Payee is one of ACME" in r.criteria and "Amount at least 2,000.00" in r.criteria


def test_columns_sort_and_top_n():
    r = rb.run(spec(columns=["payee", "amount"], sort={"by": "amount", "desc": True}, top_n=2), data=CHECKS)
    assert list(r.table.columns) == ["Payee", "Amount"] and list(r.table["Amount"]) == [3000.0, 2000.0]
    assert r.matching_rows == 5 and len(r.detail_all) == 5 and "Bank cleared" in r.detail_all.columns


def test_grouping_subtotals_and_grand_total():
    r = rb.run(spec(group_by=[{"field": "payee"}], aggregates=[{"fn": "sum", "field": "amount"}, {"fn": "count", "field": None}],
                    sort={"by": "agg:sum:amount", "desc": True}, chart="bar"), data=CHECKS)
    t = r.table
    assert list(t.columns) == ["Payee", "Sum of Amount", "Count of rows"]
    assert list(t["Payee"]) == ["ACME", "Lender Co", "S&P Tires", "Grand total"]
    assert list(t["Sum of Amount"]) == [4000.0, 2000.0, 1200.0, 7200.0] and r.row_types[-1] == "total"
    assert r.headline == ("Sum of Amount", 7200.0, "money") and r.kinds["Count of rows"] == "count"
    assert r.chart["x"] == "Payee" and len(r.chart["data"]) == 3

    two = rb.run(spec(group_by=[{"field": "track"}, {"field": "check_date", "part": "month"}],
                      aggregates=[{"fn": "sum", "field": "amount"}, {"fn": "avg", "field": "amount"}]), data=CHECKS)
    rows = list(zip(two.table["Track"], two.table["Month (Check date)"], two.table["Sum of Amount"]))
    assert rows == [("Private Lender", "Oct 2026", 2000.0), ("Private Lender", "Subtotal", 2000.0),
                    ("Supplier", "Aug 2026", 700.0), ("Supplier", "Sep 2026", 4500.0), ("Supplier", "Subtotal", 5200.0),
                    ("Grand total", "", 7200.0)]
    assert two.table["Average of Amount"].iloc[-1] == pytest.approx(7200.0 / 5)

    top = rb.run(spec(group_by=[{"field": "payee"}], aggregates=[{"fn": "sum", "field": "amount"}], sort={"by": "agg:sum:amount", "desc": True},
                      top_n=1), data=CHECKS)
    assert list(top.table["Payee"]) == ["ACME", "Grand total"] and top.table["Sum of Amount"].iloc[-1] == 4000.0


def test_exports_are_valid():
    r = rb.run(spec(group_by=[{"field": "payee"}], chart="bar"), data=CHECKS)
    assert csv_bytes(r.table, kinds=r.kinds).decode("utf-8-sig").splitlines()[0] == "Payee,Count of rows,Sum of Amount"
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes(r.table, kinds=r.kinds)))
    assert wb.active["C2"].number_format == "#,##0.00" and wb.active["B2"].number_format == "#,##0"
    pdf = rb.generate_report_pdf(r)
    doc = pdfium.PdfDocument(pdf)
    text = " ".join(" ".join(doc[i].get_textpage().get_text_range().split()) for i in range(len(doc)))
    assert "Grand total" in text and "S&P Tires" in text and "7,200.00" in text


def test_every_dataset_loads_and_runs_on_an_empty_database(db):
    for key, ds in rb.DATASETS.items():
        df = rb.load(key)
        assert list(df.columns) == list(ds.fields), key
        rb.run(rb.blank_spec(key), data=df)


def test_saved_reports_lifecycle(db):
    ok, msg, rid = rb.save_report("Checks to ACME", spec(filters=[{"field": "payee", "op": "is one of", "value": ["ACME"]}]))
    assert ok and rid
    assert not rb.save_report("checks to acme", spec())[0]                     # names are unique (any case)
    assert not rb.save_report("Broken", spec(columns=["nope"]))[0]
    ok, _, same = rb.save_report("Checks to ACME (edited)", spec(preset="This month"), report_id=rid)
    assert ok and same == rid and rb.get_saved(rid)["spec"]["preset"] == "This month"
    ok, _, dup = rb.duplicate_report(rid)
    assert ok and rb.get_saved(dup)["name"] == "Copy of Checks to ACME (edited)"
    assert rb.duplicate_report(rid)[0] and "Copy of Checks to ACME (edited) (2)" in [r["name"] for r in rb.list_saved()]
    rb.set_flags(rid, pin_dashboard=True, include_in_overview=True)
    assert [r["report_id"] for r, _ in rb.pinned_results()] == [rid]
    res = rb.overview_results(date(2026, 8, 1), date(2026, 8, 31))
    assert res[0][1].date_range == (date(2026, 8, 1), date(2026, 8, 31))       # run for the overview's own period
    rb.mark_run(rid)
    assert rb.get_saved(rid)["last_run_at"]
    rb.delete_report(dup)
    assert rb.get_saved(dup) is None
    import sqlite3
    with sqlite3.connect(db) as conn:
        acts = [r[0] for r in conn.execute("SELECT user_action FROM audit_logs WHERE entity_name = 'saved_reports'")]
    assert acts.count("REPORT_SAVED") == 4 and acts.count("REPORT_DELETED") == 1


def test_overview_includes_saved_reports(db):
    from engines.overview_pdf_service import generate_overview_pdf
    from engines.overview_report_service import get_overview

    ok, _, rid = rb.save_report("Supplier checks by payee", spec(group_by=[{"field": "payee"}]))
    rb.set_flags(rid, include_in_overview=True)
    ov = get_overview("Month", date(2026, 9, 15))
    assert [row["name"] for row, _ in ov["my_reports"]] == ["Supplier checks by payee"]
    pdf = generate_overview_pdf(ov, short=True)
    doc = pdfium.PdfDocument(pdf)
    pages = [" ".join(doc[i].get_textpage().get_text_range().split()) for i in range(len(doc))]
    mine = next(i for i, t in enumerate(pages) if "My reports Supplier checks by payee" in t) + 1
    assert f"My reports {mine}" in pages[0]
