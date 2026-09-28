"""Sales step S1 (2026-09-29): items and stock. Owner's choices: items by description with system codes; stock on hand from day one;
weighted-average cost at every stock in; a sale above stock allowed and flagged; corrections only by an owner adjustment with a
reason; the opening price list typed into a template by Sales and confirmed by the owner; goods only (tires, batteries, mags).
The stock ledger is only added to; the item's on hand always equals the sum of its movements."""
import io
from datetime import timedelta
from decimal import Decimal

import pytest

pytest.importorskip("django", reason="Django is installed only in the ERP environment (.venv)")

PASSWORD = "sandbox-Test-pass-42"
_SINCE = {"log_id": 0}


@pytest.fixture
def s1(erp):
    from django.db import connection

    from core.models import Role, User

    for role in (Role.OWNER, Role.SALES, Role.VIEWER):
        User.objects.filter(email=f"{role}@s1.example.com").delete()
    users = {role: User.objects.create_user(f"{role}@s1.example.com", f"{role.title()} S1", role, PASSWORD, must_change_password=False)
             for role in (Role.OWNER, Role.SALES, Role.VIEWER)}
    with connection.cursor() as cur:
        cur.execute("SELECT COALESCE(MAX(log_id), 0) FROM audit_logs")
        _SINCE["log_id"] = cur.fetchone()[0]
    return users


def _audit(action):
    from django.db import connection

    with connection.cursor() as cur:
        cur.execute("SELECT user_email, entity_id, old_values, new_values FROM audit_logs WHERE user_action = %s AND log_id > %s ORDER BY log_id",
                    [action, _SINCE["log_id"]])
        return cur.fetchall()


def _client(user):
    from django.test import Client

    c = Client()
    c.force_login(user)
    return c


def _supplier(name):
    """A party that is an AP supplier (the stock-in supplier list)."""
    import sqlite3

    from django.conf import settings

    from core.parties import sync_ap_links

    conn = sqlite3.connect(settings.DATABASES["default"]["NAME"])
    conn.execute("INSERT INTO creditors (creditor_name, creditor_type) VALUES (?, 'SUPPLIER')", (name,))
    conn.commit()
    conn.close()
    sync_ap_links()
    from core.models import Party

    return Party.objects.get(name=name)


def _as(user):
    from engines.acting_user import acting_as

    return acting_as(user.email)


def _ledger_ok(item):
    item.refresh_from_db()
    return sum(item.moves.values_list("qty", flat=True)) == item.on_hand


# ---- items --------------------------------------------------------------------------------------------------------------------------

def test_items_get_codes_per_category_and_the_same_item_cannot_be_added_twice(s1):
    from inventory import stock

    with _as(s1["sales"]):
        a = stock.add_item("tire", "CodeTest", "185/65R15")
        b = stock.add_item("tire", "CodeTest", "185/65R15", "old stock 2023")          # a remark tells two lots apart
        c = stock.add_item("battery", "CodeTest", "NS60")
        with pytest.raises(stock.StockError, match="already on the list"):
            stock.add_item("tire", "  codetest ", "185/65r15.")
    assert a.code.startswith("TR-") and b.code.startswith("TR-") and int(b.code[3:]) == int(a.code[3:]) + 1
    assert c.code.startswith("BA-")
    assert _audit("INV_ITEM_ADDED")[0][0] == "sales@s1.example.com"


# ---- stock in and weighted average ---------------------------------------------------------------------------------------------------

def test_stock_in_updates_on_hand_and_the_weighted_average(s1):
    from engines.clock import manila_today
    from inventory import stock

    sup = _supplier("Avg Supplier S1")
    with _as(s1["sales"]):
        item = stock.add_item("tire", "AvgBrand", "185/65R15")
        stock.record_stock_in(sup, manila_today(), "DR 1", [(item, 12, "3150")])
        preview = stock.check_stock_in(sup, manila_today(), [(stock.Item.objects.get(pk=item.pk), 10, "3,300.00")])
        assert preview[0]["avg_after"].quantize(Decimal("0.01")) == Decimal("3218.18")
        doc = stock.record_stock_in(sup, manila_today(), "DR 2", [(item, 10, "3300")])
    item.refresh_from_db()
    assert item.on_hand == 22 and item.avg_cost.quantize(Decimal("0.01")) == Decimal("3218.18")
    assert doc.number.endswith(" - RR " + doc.number[-4:]) and doc.total == Decimal("33000.00")
    assert _ledger_ok(item) and len(_audit("INV_STOCK_IN")) == 2


def test_a_delivery_after_the_stock_went_below_zero_sets_the_average_to_its_cost(s1):
    from engines.clock import manila_today
    from inventory import stock
    from inventory.models import MoveKind

    sup = _supplier("Negative Supplier S1")
    with _as(s1["owner"]):
        item = stock.add_item("tire", "NegBrand", "195R15C")
        stock.record_stock_in(sup, manila_today(), "", [(item, 2, "4000")])
        from django.db import transaction

        with transaction.atomic():                                        # what an S2 sale of 3 will do: allowed, stock -1
            stock._move(stock.Item.objects.get(pk=item.pk), MoveKind.SALE, -3, Decimal("4000"), manila_today())
        item.refresh_from_db()
        assert item.on_hand == -1 and item.is_negative
        stock.record_stock_in(sup, manila_today(), "", [(item, 4, "4500")])
    item.refresh_from_db()
    assert item.on_hand == 3 and item.avg_cost == Decimal("4500") and _ledger_ok(item)


@pytest.mark.parametrize("lines, message", [
    ([], "at least one item"),
    ([("ITEM", 0, "100")], "at least 1"),
    ([("ITEM", "1.5", "100")], "whole pieces"),
    ([("ITEM", 1, "-5")], "can't be negative"),
    ([("ITEM", 1, "abc")], "not an amount"),
    ([("ITEM", 1, "100"), ("ITEM", 2, "100")], "on two lines"),
])
def test_stock_in_mistakes_are_refused(s1, lines, message):
    from engines.clock import manila_today
    from inventory import stock

    sup = _supplier(f"Refuse Supplier S1 {message}")
    with _as(s1["owner"]):
        item = stock.add_item("mags", f"RefuseBrand {message}", "15 4x100")
        with pytest.raises(stock.StockError, match=message):
            stock.record_stock_in(sup, manila_today(), "", [(item if x == "ITEM" else x, q, c) for x, q, c in lines])
    item.refresh_from_db()
    assert item.on_hand == 0 and not item.moves.exists()


def test_future_dates_and_a_missing_supplier_are_refused(s1):
    from engines.clock import manila_today
    from inventory import stock

    with _as(s1["owner"]):
        item = stock.add_item("tire", "FutureBrand", "205/55R16")
        with pytest.raises(stock.StockError, match="future"):
            stock.record_stock_in(_supplier("Future Supplier S1"), manila_today() + timedelta(days=1), "", [(item, 1, "100")])
        with pytest.raises(stock.StockError, match="supplier"):
            stock.record_stock_in(None, manila_today(), "", [(item, 1, "100")])


def test_the_stock_in_page_checks_first_then_saves(s1):
    from inventory import stock
    from inventory.models import StockIn

    sup = _supplier("Page Supplier S1")
    with _as(s1["owner"]):
        item = stock.add_item("tire", "PageBrand", "175/70R13")
    sales = _client(s1["sales"])
    data = {"supplier": sup.pk, "received_on": sales.get("/stock-in/new/").context["today"], "supplier_ref": "DR 77",
            "item_0": f"{item.code} · {item.description}", "qty_0": "5", "cost_0": "2,500.00", "action": "check"}
    resp = sales.post("/stock-in/new/", data)
    assert resp.status_code == 200 and resp.context["preview"][0]["on_hand_after"] == 5
    assert not StockIn.objects.filter(supplier_ref="DR 77").exists()          # nothing saved by Check
    data["action"] = "save"
    resp = sales.post("/stock-in/new/", data)
    doc = StockIn.objects.get(supplier_ref="DR 77")
    assert resp["Location"] == f"/stock-in/{doc.pk}/"
    item.refresh_from_db()
    assert item.on_hand == 5 and item.avg_cost == Decimal("2500")


def test_owner_reverses_a_stock_in_and_the_average_goes_back(s1):
    from engines.clock import manila_today
    from inventory import stock

    sup = _supplier("Reverse Supplier S1")
    with _as(s1["owner"]):
        item = stock.add_item("tire", "RevBrand", "185R14C")
        stock.record_stock_in(sup, manila_today(), "", [(item, 10, "3000")])
        wrong = stock.record_stock_in(sup, manila_today(), "", [(item, 10, "4000")])
    assert _client(s1["sales"]).post(f"/stock-in/{wrong.pk}/reverse/", {"reason": "twice"}).status_code == 403
    owner = _client(s1["owner"])
    owner.post(f"/stock-in/{wrong.pk}/reverse/", {"reason": ""})
    wrong.refresh_from_db()
    assert wrong.reversed_at is None                                          # a reason is required
    owner.post(f"/stock-in/{wrong.pk}/reverse/", {"reason": "entered the wrong cost"})
    item.refresh_from_db()
    wrong.refresh_from_db()
    assert wrong.reversed_at and item.on_hand == 10 and item.avg_cost == Decimal("3000") and _ledger_ok(item)
    assert _audit("INV_STOCK_IN_REVERSED")
    with pytest.raises(stock.StockError, match="already reversed"):
        stock.reverse_stock_in(wrong, "again")


# ---- owner adjustment ---------------------------------------------------------------------------------------------------------------

def test_only_the_owner_adjusts_to_a_count_with_a_reason(s1):
    from engines.clock import manila_today
    from inventory import stock

    sup = _supplier("Adjust Supplier S1")
    with _as(s1["owner"]):
        item = stock.add_item("battery", "AdjBrand", "NS40")
        stock.record_stock_in(sup, manila_today(), "", [(item, 12, "3000")])
    assert _client(s1["sales"]).post(f"/items/{item.pk}/adjust/", {"counted": "10", "reason": "count"}).status_code == 403
    owner = _client(s1["owner"])
    owner.post(f"/items/{item.pk}/adjust/", {"counted": "10", "reason": ""})
    item.refresh_from_db()
    assert item.on_hand == 12                                                 # no reason, no change
    owner.post(f"/items/{item.pk}/adjust/", {"counted": "10", "reason": "count Sep 30"})
    item.refresh_from_db()
    assert item.on_hand == 10 and item.avg_cost == Decimal("3000") and _ledger_ok(item)
    move = item.moves.first()
    assert move.kind == "adjustment" and move.qty == -2 and move.reason == "count Sep 30"
    assert _audit("INV_ADJUSTED")[0][0] == "owner@s1.example.com"
    with pytest.raises(stock.StockError, match="below zero"):
        stock.adjust_to_count(item, -1, "oops")


def test_items_below_zero_are_flagged_for_the_owner(s1):
    from django.db import transaction

    from engines.clock import manila_today
    from inventory import stock
    from inventory.models import MoveKind

    with _as(s1["owner"]):
        item = stock.add_item("tire", "FlagBrand", "215/65R16")
        with transaction.atomic():
            stock._move(stock.Item.objects.get(pk=item.pk), MoveKind.SALE, -1, Decimal("0"), manila_today())
    html = _client(s1["owner"]).get("/items/").content.decode()
    assert "less than zero on hand" in html
    assert _client(s1["owner"]).get("/items/", {"show": "low", "q": "FlagBrand"}).context["rows"][0]["i"] == stock.Item.objects.get(pk=item.pk)


# ---- opening list -------------------------------------------------------------------------------------------------------------------

def _filled_template(rows_by_sheet):
    from openpyxl import load_workbook

    from inventory import opening

    wb = load_workbook(io.BytesIO(opening.build_template()))
    for sheet, rows in rows_by_sheet.items():
        ws = wb[sheet]
        for r, values in enumerate(rows, start=2):
            for c, v in enumerate(values, start=1):
                ws.cell(row=r, column=c, value=v)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    out.name = "price_list.xlsx"
    return out


def test_the_template_has_the_paper_lists_columns():
    pytest.importorskip("django")
    from openpyxl import load_workbook

    from inventory import opening

    wb = load_workbook(io.BytesIO(opening.build_template()))
    assert wb.sheetnames == ["How to fill", "Tires", "Batteries", "Mags"]
    assert [c.value for c in wb["Tires"][1]] == ["Brand", "Size", "Qty", "Remarks", "Cost"]
    assert [c.value for c in wb["Batteries"][1]] == ["Brand", "Model", "Qty", "Remarks", "Cost"]


def test_opening_list_mistakes_are_listed_and_nothing_is_saved(s1):
    from inventory.models import OpeningImport

    before = OpeningImport.objects.count()
    f = _filled_template({"Tires": [("OpenErr", "", 2, "", 100), ("OpenErr", "185/65R15", "x", "", 100),
                                    ("OpenErr", "195R15C", 1.5, "", 100), ("OpenErr", "205R16", 1, "", None),
                                    ("OpenErr", "215R16", 1, "", 100), ("openerr", "215r16", 3, "", 100)]})
    resp = _client(s1["sales"]).post("/opening-list/", {"file": f})
    errors = resp.context["result"]["errors"]
    assert len(errors) == 5 and OpeningImport.objects.count() == before
    text = " ".join(errors)
    assert "brand and size" in text and "numbers" in text and "whole pieces" in text and "cost is missing" in text and "same item" in text
    bad = io.BytesIO(b"not excel")
    bad.name = "x.xlsx"
    assert "can read" in _client(s1["sales"]).post("/opening-list/", {"file": bad}).context["result"]["errors"][0]


def test_sales_uploads_the_opening_list_and_the_owner_confirms_the_totals(s1):
    from inventory.models import ImportStatus, Item, OpeningImport

    for imp in OpeningImport.objects.filter(status=ImportStatus.PENDING):  # a clean start for this test
        imp.status = ImportStatus.DISCARDED
        imp.save()
    f = _filled_template({"Tires": [("OpenOK", "185/65R15", 12, "", "3,150.00"), ("OpenOK", "185/65R15", 2, "old stock 2023", 2700)],
                          "Batteries": [("OpenOK", "NS60", 5, "", 4050)], "Mags": [("OpenOK", "15 4x100", 0, "set of 4", 9800)]})
    sales = _client(s1["sales"])
    sales.post("/opening-list/", {"file": f})
    imp = OpeningImport.objects.get(status=ImportStatus.PENDING)
    assert (imp.items_count, imp.pieces, imp.value) == (4, 19, Decimal("63450.00"))
    assert imp.uploaded_by == "sales@s1.example.com" and _audit("INV_OPENING_UPLOADED")
    assert sales.post(f"/opening-list/{imp.pk}/decide/", {"decision": "confirm"}).status_code == 403
    f2 = _filled_template({"Tires": [("OpenOK2", "175R13", 1, "", 100)]})
    assert "waiting for the owner" in sales.post("/opening-list/", {"file": f2}, follow=True).content.decode()

    _client(s1["owner"]).post(f"/opening-list/{imp.pk}/decide/", {"decision": "confirm"})
    imp.refresh_from_db()
    assert imp.status == ImportStatus.CONFIRMED and imp.decided_by == "owner@s1.example.com"
    tire = Item.objects.get(brand="OpenOK", size_model="185/65R15", remark="")
    mags = Item.objects.get(brand="OpenOK", category="mags")
    assert tire.on_hand == 12 and tire.avg_cost == Decimal("3150") and tire.moves.get().kind == "opening"
    assert mags.on_hand == 0 and mags.avg_cost == Decimal("9800") and not mags.moves.exists()
    assert _audit("INV_OPENING_CONFIRMED")

    again = _filled_template({"Tires": [("OpenOK", "185/65R15", 99, "", 1)]})
    errors = sales.post("/opening-list/", {"file": again}).context["result"]["errors"]
    assert "already has stock movements" in errors[0]


def test_the_owner_can_discard_a_waiting_list(s1):
    from inventory.models import ImportStatus, Item, OpeningImport

    for imp in OpeningImport.objects.filter(status=ImportStatus.PENDING):
        imp.status = ImportStatus.DISCARDED
        imp.save()
    _client(s1["sales"]).post("/opening-list/", {"file": _filled_template({"Tires": [("Discarded", "185R14", 3, "", 100)]})})
    imp = OpeningImport.objects.get(status=ImportStatus.PENDING)
    _client(s1["owner"]).post(f"/opening-list/{imp.pk}/decide/", {"decision": "discard", "reason": "typos"})
    imp.refresh_from_db()
    assert imp.status == ImportStatus.DISCARDED and not Item.objects.filter(brand="Discarded").exists()


# ---- who sees what ------------------------------------------------------------------------------------------------------------------

def test_every_inventory_page_declares_its_roles_and_viewers_are_refused(s1):
    from inventory import urls

    viewer = _client(s1["viewer"])
    for pattern in urls.urlpatterns:
        assert getattr(pattern.callback, "allowed_roles", None), pattern.name
    for path in ("/items/", "/items/add/", "/stock-in/", "/stock-in/new/", "/opening-list/", "/opening-list/template/"):
        assert viewer.get(path).status_code == 403, path


def test_sales_sees_items_stock_in_and_the_template(s1):
    sales = _client(s1["sales"])
    for path in ("/items/", "/items/add/", "/stock-in/", "/stock-in/new/", "/opening-list/"):
        assert sales.get(path).status_code == 200, path
    t = sales.get("/opening-list/template/")
    assert t.status_code == 200 and "spreadsheetml" in t["Content-Type"]
    home = sales.get("/").content.decode()
    assert "/items/" in home and "/stock-in/" in home


def test_stock_in_offers_only_suppliers_active_in_ap(s1):
    import sqlite3

    from django.conf import settings

    active, inactive = _supplier("Active Supplier S1"), _supplier("Inactive Supplier S1")
    conn = sqlite3.connect(settings.DATABASES["default"]["NAME"])
    conn.execute("UPDATE creditors SET is_active = 0 WHERE creditor_name = 'Inactive Supplier S1'")
    conn.commit()
    conn.close()
    offered = list(_client(s1["sales"]).get("/stock-in/new/").context["suppliers"])
    assert active in offered and inactive not in offered
