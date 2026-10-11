"""The standard price template (owner, 2026-10-10: "standardize the fields for each type tires, battery, rims, mags"; "a drop down instead of just
data validation"; no AI, no file reading): one sheet per type with its own fixed fields, real drop-downs, pre-filled per supplier; tires saved earlier
with their ply and terrain as free text are moved to the fields once, so the same tire is found again."""
import io
from datetime import date

from openpyxl import load_workbook

from conftest import fetchone, make_creditor
from engines import price_template_service as pt
from engines import product_catalog as pc


def _fill(data: bytes, sheet: str, rows: list[dict]) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    ws = wb[sheet]
    head = {" ".join(str(c.value or "").replace("*", "").split()).lower(): c.column for c in ws[1]}
    r = ws.max_row + 1 if ws.cell(row=2, column=1).value else 2
    while ws.cell(row=r, column=head["brand"]).value:
        r += 1
    r = next(i for i in range(2, 5000) if not ws.cell(row=i, column=head["brand"]).value)
    for row in rows:
        for k, v in row.items():
            ws.cell(row=r, column=head[k.lower()], value=v)
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_one_sheet_per_type_with_its_own_fields_and_drop_downs(db):
    wb = load_workbook(io.BytesIO(pt.build_type_template()))
    assert wb.sheetnames == ["How to fill", "Lists", "Tires", "Battery", "Mags", "Rims"] and wb["Lists"].sheet_state == "hidden"
    tires = [str(c.value).replace(" *", "") for c in wb["Tires"][1]]
    assert tires[:7] == ["Brand", "Size", "Ply / load-speed", "Terrain", "Pattern", "Type", "Tube"]
    assert [str(c.value).replace(" *", "") for c in wb["Battery"][1]][:7] == ["Brand", "Model / group", "Volts", "Ah", "CCA", "Type", "Terminal"]
    assert [str(c.value).replace(" *", "") for c in wb["Rims"][1]][:8] == ["Brand", "Size", "Holes", "PCD", "Offset", "Tube", "Disc thickness", "Color"]
    assert [str(c.value).replace(" *", "") for c in wb["Mags"][1]][:7] == ["Brand", "Model", "Size", "PCD", "Offset", "Center bore", "Finish / color"]
    lists = [dv for dv in wb["Tires"].data_validations.dataValidation if dv.type == "list"]
    assert len(lists) == 4                                                          # Brand, Terrain, Type, Tube
    cols = {wb["Lists"].cell(row=1, column=j).value: [wb["Lists"].cell(row=i, column=j).value for i in range(2, 12) if wb["Lists"].cell(row=i, column=j).value]
            for j in range(1, wb["Lists"].max_column + 1)}
    assert cols["Tires.terrain"] == ["A/T", "M/T", "H/T", "R/T", "Lug", "Rib", "Mix"] and cols["Battery.terminal"] == ["L", "R"]
    assert all(dv.showDropDown in (None, False) for dv in lists)           # the arrow shows (showDropDown=True would HIDE it)


def test_a_filled_sheet_reads_into_the_fields_and_a_bad_choice_is_named_by_row(db):
    cid = make_creditor(db, "New Century City Marketing Corp.")
    data = _fill(pt.build_type_template(cid), "Tires", [
        {"Brand": "Westlake", "Size": "315/80R22.5", "Ply / load-speed": "20PR", "Terrain": "Lug", "Pattern": "MD777", "Type": "TBR", "List price": 16800},
        {"Brand": "Westlake", "Size": "315/80R22.5", "Ply / load-speed": "20PR", "Terrain": "Snow", "Pattern": "CR926", "List price": 15700},
        {"Brand": "Westlake", "Size": "12.00-20", "Ply / load-speed": "18PR", "Terrain": "lug", "Pattern": "CL991", "List price": "abc"}])
    parsed = pt.parse_workbook(data, default_date=date(2026, 9, 15))
    assert parsed["suppliers"] == ["New Century City Marketing Corp."]
    assert len(parsed["rows"]) == 1 and parsed["rows"][0]["values"]["terrain"] == "Lug" and parsed["rows"][0]["values"]["ply"] == "20PR"
    assert any("Tires row 3" in p and "Terrain 'Snow' is not one of" in p for p in parsed["problems"])
    assert any("Tires row 4" in p and "List price 'abc' is not a number" in p for p in parsed["problems"])
    ok, msg, counts = pt.import_workbook(cid, parsed)
    assert ok and counts["added"] == 1 and counts["products"] == 1


def test_prefilled_round_trip_one_changed_price(db):
    cid = make_creditor(db, "New Century City Marketing Corp.")
    first = _fill(pt.build_type_template(cid), "Tires", [
        {"Brand": "Westlake", "Size": "315/80R22.5", "Ply / load-speed": "20PR", "Terrain": "Lug", "Pattern": "MD777", "List price": 16800},
        {"Brand": "Westlake", "Size": "315/80R22.5", "Ply / load-speed": "20PR", "Terrain": "Mix", "Pattern": "CR926", "List price": 15700, "Remarks": "Mining"}])
    assert pt.import_workbook(cid, pt.parse_workbook(first, default_date=date(2026, 9, 15)))[0]
    wb = load_workbook(io.BytesIO(pt.build_type_template(cid, on=date(2026, 10, 1))))
    ws = wb["Tires"]
    head = {str(c.value).replace(" *", ""): c.column for c in ws[1]}
    rows = [[ws.cell(row=r, column=head[k]).value for k in ("Brand", "Size", "Ply / load-speed", "Terrain", "Pattern", "List price", "Remarks")] for r in (2, 3)]
    assert ["Westlake", "315/80R22.5", "20PR", "Lug", "MD777", 16800, None] in rows and ["Westlake", "315/80R22.5", "20PR", "Mix", "CR926", 15700, "Mining"] in rows
    assert str(ws.cell(row=2, column=head["Changed"]).value).startswith("=IF(") and ws.column_dimensions[ws.cell(row=1, column=head["Old price"]).column_letter].hidden
    target = next(r for r in (2, 3) if ws.cell(row=r, column=head["Pattern"]).value == "MD777")
    ws.cell(row=target, column=head["List price"], value=17500)
    buf = io.BytesIO()
    wb.save(buf)
    parsed = pt.parse_workbook(buf.getvalue(), default_date=date(2026, 10, 1))
    assert parsed["problems"] == []
    cmp = pt.compare(cid, parsed)
    assert sorted(cmp["status"]) == ["changed", "same"]
    ok, msg, counts = pt.import_workbook(cid, parsed)
    assert ok and counts["added"] == 1 and counts["products"] == 0                      # the same two items; one new price


def test_tires_saved_with_free_text_others_are_moved_to_the_fields_once(db):
    cid = make_creditor(db, "New Century City Marketing Corp.")
    pid, name, _ = pc.get_or_create_product("Tires", "Apollo", size="12.00R20", pattern="EnduTrax MA", remark="18PR Mix")     # as the one-sheet template saved it
    pid2, _, _ = pc.get_or_create_product("Tires", "Apollo", size="12.00-20", pattern="MINE LUG", remark="20PR Mining Lug")
    res = pc.refield_tires()
    assert res["changed"] == 2 and res["duplicates"] == []
    assert pc.refield_tires(apply=True)["changed"] == 2 and pc.refield_tires()["changed"] == 0     # once
    assert fetchone(db, "SELECT name, remark FROM items WHERE id = ?", (pid,)) == (name, "")       # the name stays; nothing left in the remark
    assert fetchone(db, "SELECT remark FROM items WHERE id = ?", (pid2,))[0] == "Mining"
    again, _, created = pc.get_or_create_product("Tires", "Apollo", size="12.00R20", ply="18PR", terrain="Mix", pattern="EnduTrax MA")
    assert again == pid and not created                                                        # the per-type row finds the same tire
    assert fetchone(db, "SELECT COUNT(*) FROM audit_logs WHERE user_action = 'ITEMS_REFIELDED'")[0] == 1


def test_the_old_single_sheet_and_the_new_sheets_give_one_tire_one_key(db):
    vals, remark, bad = pc.tidy_tire("MD777", "315/80R22.5-20PR", "", "Lug TL")
    vals, remark = pc.field_tire_remark(vals, remark)
    assert bad == [] and (vals["size"], vals["ply"], vals["terrain"], vals["tube"], remark) == ("315/80R22.5", "20PR", "Lug", "TL", "")
    assert pc.product_key("Tires", "Westlake", remark, **vals) == pc.product_key("Tires", "Westlake", size="315/80R22.5", ply="20PR", terrain="Lug", pattern="MD777", tube="TL")


def test_cells_copied_as_the_sunny_pcr_list_prints_them_land_in_their_own_fields(db):
    # owner, 2026-10-10: "SUNNY PCR (09-15-2026).pdf -- consider this in the template": Rim | Size | LI/SI | Pattern | Price
    cid = make_creditor(db, "Sunny Tires")
    data = _fill(pt.build_type_template(cid), "Tires", [
        {"Brand": "Sunny", "Size": "285/45R22", "Ply / load-speed": "114V XL", "Pattern": "NU025 H/T", "List price": 6200},
        {"Brand": "Sunny", "Size": "LT255/55R19-10PR", "Ply / load-speed": "111T", "Pattern": "SU011 A/T", "List price": 8000},
        {"Brand": "Sunny", "Size": "215/70R16C6", "Ply / load-speed": "108/106T", "Pattern": "NL106", "Type": "LT", "List price": 3900},
        {"Brand": "Sunny", "Size": "265/50R20", "Ply / load-speed": "111V XL", "Pattern": "NU006 / NU033 A/T", "List price": 6000}])
    parsed = pt.parse_workbook(data, default_date=date(2026, 9, 15))
    rows = {r["values"]["pattern"]: r for r in parsed["rows"]}
    a, b, c = rows["NU025"], rows["SU011"], rows["NL106"]
    assert (a["values"]["size"], a["values"]["ply"], a["values"]["terrain"], a["item_remark"]) == ("285/45R22", "114V", "H/T", "XL")
    assert (b["values"]["size"], b["values"]["ply"], b["values"]["terrain"]) == ("LT255/55R19", "10PR 111T", "A/T")
    assert (c["values"]["size"], c["values"]["ply"], c["values"]["type"], c["item_remark"]) == ("215/70R16", "108/106T", "LT", "C6")
    assert len(parsed["rows"]) == 3 and any("Tires row 5" in p and "two patterns" in p for p in parsed["problems"])


def test_the_full_words_are_accepted_where_the_short_codes_live(db):
    # owner, 2026-10-11: his list said "Tubeless" and "Highway-Terrain" and 89 rows were refused -- the obvious spellings count
    cid = make_creditor(db, "New Century City Marketing Corp.")
    data = _fill(pt.build_type_template(cid), "Tires", [
        {"Brand": "Westlake", "Size": "265/65R17", "Ply / load-speed": "112T", "Terrain": "Highway-Terrain", "Pattern": "SL369", "Type": "PCR", "Tube": "Tubeless", "List price": 7400},
        {"Brand": "Westlake", "Size": "31X10.50R15", "Ply / load-speed": "109S", "Terrain": "All Terrain", "Pattern": "SL309", "Tube": "Tube type", "List price": 8200},
        {"Brand": "Westlake", "Size": "12.00-20", "Ply / load-speed": "18PR", "Terrain": "Mud-Terrain", "Pattern": "CL991", "List price": 15700}])
    parsed = pt.parse_workbook(data, default_date=date(2026, 9, 15))
    assert parsed["problems"] == []
    got = {r["values"]["pattern"]: (r["values"]["terrain"], r["values"].get("tube", "")) for r in parsed["rows"]}
    assert got["SL369"] == ("H/T", "TL") and got["SL309"] == ("A/T", "TT") and got["CL991"][0] == "M/T"
    assert pc.terrain_of("Rugged-Terrain") == "R/T" and pc.tube_of("TUBELESS") == "TL" and pc.tube_of("nonsense") is None
