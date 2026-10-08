"""The one standard price-list sheet (owner, 2026-10-08): Supplier, Category, Brand, Model/Pattern, Size, Terrain, Other specs, Unit, Supplier item code, List price,
Effective date, for every product. Tires: brand, size, pattern, terrain and the price are kept clean; everything else is "the others", free text that still tells two
tires apart. Older files still work."""
import io
from datetime import date

from openpyxl import load_workbook

from conftest import fetchone, make_creditor
from engines import price_template_service as pt
from engines import product_catalog as pc
from engines.price_list_service import price_on

DAY = date(2026, 9, 15)
H = ["Supplier", "Category", "Brand", "Model/Pattern", "Size", "Terrain", "Other specs", "Unit", "Supplier item code", "List price", "Effective date"]
#        Category   Brand       Model    Size          Terrain Other specs             Unit  Code    Price    Date
GOOD = [["Tires", "Westlake", "SL369", "285/50R20", "A/T", "", "pc", "WL-1", 7600, DAY],
        ["tire", "Westlake", "SA57", "305/35R24", "", "", "pc", "", "8,400", "2026-09-15"],
        ["Battery", "Motolite", "NS60", "", "", "12V 45Ah L", "pc", "", 4200, DAY],
        ["Mags", "Enkei", "RPF1", "17x7.5", "", "5x114.3 ET40 Silver", "pc", "", 9800, DAY],
        ["Rims", "Westlake", "", "8.25x22.5", "", "10H TL", "pc", "", 4150, DAY]]


def _file(rows, creditor_id=None, supplier=None) -> bytes:
    """GOOD-style rows (no Supplier) written from the Category column on; `supplier` in the Supplier column of every row."""
    wb = load_workbook(io.BytesIO(pt.build_template(creditor_id)))
    ws = wb[pt.SHEET]
    start = 2 if ws.cell(row=2, column=2).value is None else ws.max_row + 1     # a supplier template may carry only the supplier's name on row 2
    for r, vals in enumerate(rows, start=start):
        if supplier:
            ws.cell(row=r, column=1, value=supplier)
        for j, v in enumerate(vals, start=2):
            ws.cell(row=r, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _tire(model, size, terrain="", other="", price=5000):
    return ["Tires", "Westlake", model, size, terrain, other, "pc", "", price, DAY]


def test_the_template_is_one_sheet_with_the_standard_columns():
    wb = load_workbook(io.BytesIO(pt.build_template()))
    assert wb.sheetnames == ["How to fill", "Price list"]
    assert [str(c.value).replace(" *", "") for c in wb["Price list"][1]] == H
    assert any("INSTRUCTIONS FOR THE AI" in str(c.value) for c in wb["How to fill"]["A"])


def test_rows_are_read_into_clean_products():
    parsed = pt.parse_workbook(_file(GOOD))
    assert parsed["problems"] == [] and len(parsed["rows"]) == 5
    names = {r["name"] for r in parsed["rows"]}
    assert {"Westlake 285/50R20 (A/T) SL369", "Westlake 305/35R24 SA57", "Motolite NS60 12V 45Ah L", "Enkei RPF1 17x7.5 5x114.3 ET40 Silver",
            "Westlake 8.25x22.5 10H Tubeless"} == names
    bat = next(r for r in parsed["rows"] if r["category"] == "Battery")
    assert bat["values"]["ah"] == "45" and bat["values"]["volts"] == "12" and bat["values"]["terminal"] == "L"


def test_a_tire_keeps_its_crucial_fields_clean_and_the_rest_as_others():
    vals, others, bad = pc.tidy_tire("NU025 H/T", "LT255/55R19-10PR", "", "111T")
    assert bad == [] and vals == {"size": "LT255/55R19", "pattern": "NU025", "terrain": "H/T"} and others == "10PR 111T"
    vals, others, bad = pc.tidy_tire("NL106", "225/55ZR17C8", "", "109/107T")
    assert bad == [] and vals["size"] == "225/55ZR17" and others == "C8 109/107T"
    for size in ("11R22.5", "7.00R16", "31X10.50R15", "195R15", "23.5-25", "9.5R17.5", "12.00R20", "LT205/85R16", "4.00-8", "6.40/6.50-13"):
        assert pc.tidy_tire("X1", size, "", "")[2] == [], size
    assert "has two patterns" in pc.tidy_tire("NU006 / NU033", "265/50R20", "A/T", "")[2][0]
    assert "does not look like a tire size" in pc.tidy_tire("X1", "big one", "", "")[2][0]
    assert "is not A/T, H/T, M/T or R/T" in pc.tidy_tire("X1", "285/50R20", "Snow", "")[2][0]
    assert pc.tidy_tire("X1", "285/50R20", "at", "")[0]["terrain"] == "A/T"


def test_each_tire_is_unique_by_all_its_specs_and_the_order_of_the_others_does_not_matter():
    tt = pc.product_key("Tires", "Westlake", "14PR Rib TT", size="7.00R16", pattern="EZ525")
    tl = pc.product_key("Tires", "Westlake", "14PR Rib TL", size="7.00R16", pattern="EZ525")
    assert tt != tl
    assert pc.product_key("Tires", "Roadone", "Lug Mining", size="11R22.5", pattern="X") == pc.product_key("Tires", "roadone", "mining  LUG", size="11R22.5", pattern="X")


def test_tubetype_and_tubeless_get_their_own_prices(db):
    """The New Century EZ525 mix-up (2026-10-08): two tires that differ only in TT / TL each keep their own price."""
    cid = make_creditor(db, name="New Century")
    rows = [_tire("EZ525", "7.00R16", "", "14PR Rib TT", 5500), _tire("EZ525", "7.00R16", "", "14PR Rib TL", 5200)]
    ok, msg, counts = pt.import_workbook(cid, pt.parse_workbook(_file(rows)))
    assert ok and counts["added"] == 2 and counts["products"] == 2, msg
    names = [r[0] for r in __import__("sqlite3").connect(db).execute("SELECT name FROM items ORDER BY code")]
    assert names == ["Westlake 7.00R16 EZ525 14PR Rib TT", "Westlake 7.00R16 EZ525 14PR Rib TL"]
    assert price_on(cid, names[0], DAY)["unit_price"] == 5500 and price_on(cid, names[1], DAY)["unit_price"] == 5200


def test_two_items_never_share_a_name(db):
    a = pc.get_or_create_product("Tires", "Westlake", size="7.00R16", pattern="EZ525", type="LT")
    b = pc.get_or_create_product("Tires", "Westlake", size="7.00R16", pattern="EZ525", type="PCR")      # type is not in the name: the code is added
    assert a[0] != b[0] and a[1] != b[1] and b[1].endswith("[TR-0002]")


def test_the_same_product_is_found_ignoring_spaces_dashes_and_capitals():
    parsed = pt.parse_workbook(_file([GOOD[0], ["TIRES", "west-lake", "sl 369", "285/50 R20", "at", "", "pc", "", 7700, DAY]]))
    assert parsed["problems"] == ["Price list row 3: the same product and date as row 2, but a different price (7,600.00 there): keep the right one."]


def test_an_exact_repeated_row_is_skipped_with_a_note_not_a_problem():
    parsed = pt.parse_workbook(_file([GOOD[0], GOOD[2], GOOD[0]]))                 # an AI-filled file wrote the same line twice
    assert parsed["problems"] == [] and len(parsed["rows"]) == 2
    assert parsed["notes"] == ["Price list row 4: the same as row 2 (same product, date and price), so it is skipped."]


def test_sizes_with_specs_inside_are_tidied_and_reported():
    parsed = pt.parse_workbook(_file([_tire("SU011", "LT255/55R19-10PR", "A/T", "111T")]))
    assert parsed["problems"] == [] and parsed["rows"][0]["values"]["size"] == "LT255/55R19" and parsed["rows"][0]["item_remark"] == "10PR 111T"
    assert parsed["tidied"] and "LT255/55R19-10PR" in parsed["tidied"][0]


def test_problems_are_named_by_row():
    bad = [["Shoes", "Nike", "X", "1", "", "", "", "", 5, DAY],
           ["Battery", "Motolite", "NS60", "", "", "12V L", "", "", 5, DAY],
           ["Mags", "Enkei", "RPF1", "17x7.5", "", "Silver", "", "", 5, DAY],
           ["Tires", "", "SL369", "285/50R20", "", "", "", "", "call", DAY]]
    text = "\n".join(pt.parse_workbook(_file(bad))["problems"])
    for needle in ("Price list row 2: Category 'Shoes'", "Price list row 4: PCD is empty", "Price list row 5: Brand is empty; List price 'call'"):
        assert needle in text, needle
    assert "row 3" not in text                                              # a battery without its Ah is accepted (owner, 2026-10-08)


def test_import_then_the_prefilled_download_comes_back_the_same(db):
    cid = make_creditor(db, name="New Century")
    ok, msg, counts = pt.import_workbook(cid, pt.parse_workbook(_file(GOOD + [_tire("EZ525", "7.00R16", "", "14PR Rib TT")])))
    assert ok and counts["added"] == 6 and counts["products"] == 6
    assert price_on(cid, "Motolite NS60 12V 45Ah L", date(2026, 10, 1))["unit_price"] == 4200
    again = pt.parse_workbook(pt.build_template(cid, on=date(2026, 10, 1)))
    assert again["problems"] == [] and len(again["rows"]) == 6 and set(pt.compare(cid, again)["status"]) == {"same"}


def test_an_older_file_with_a_sheet_per_category_still_imports(db):
    cid = make_creditor(db, name="New Century")
    wb = load_workbook(io.BytesIO(pt.build_template(layout="by_category")))
    ws = wb["Battery"]
    header = [str(c.value).replace(" *", "") for c in ws[1]]
    for k, v in {"Brand": "Motolite", "Model / group": "NS60", "Volts": 12, "Ah": 45, "Terminal": "L", "List price": 4200, "Effective date": DAY}.items():
        ws.cell(row=2, column=header.index(k) + 1, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    old = pt.parse_workbook(buf.getvalue())
    assert old["problems"] == [] and pt.import_workbook(cid, old)[0]
    new = pt.parse_workbook(_file([GOOD[2]]))
    assert set(pt.compare(cid, new)["status"]) == {"same"}                       # the same product, whichever file it came from
    assert fetchone(db, "SELECT COUNT(*) FROM items")[0] == 1


def test_rekey_recomputes_keys_with_every_spec_and_never_merges(db):
    from engines.common import write_tx
    cid = make_creditor(db, name="New Century")
    pt.import_workbook(cid, pt.parse_workbook(_file(GOOD)))
    with write_tx() as conn:                                                   # as an item saved before every spec counted
        conn.execute("UPDATE items SET key = 'battery|motolite|ns60|l' WHERE brand = 'Motolite'")
    res = pc.rekey_items()
    assert res == {"changed": 1, "duplicates": []}
    assert fetchone(db, "SELECT key FROM items WHERE brand = 'Motolite'")[0] == "battery|motolite|ns60|l"   # only a report without --apply
    assert pc.rekey_items(apply=True)["changed"] == 1 and pc.rekey_items()["changed"] == 0
    assert "45" in fetchone(db, "SELECT key FROM items WHERE brand = 'Motolite'")[0]


def test_the_supplier_in_the_file_must_be_the_one_chosen(db):
    nc = make_creditor(db, name="New Century City Marketing Corp.")
    amk = make_creditor(db, name="AMK Auto Parts and Accessories Trading")
    assert pt.same_supplier("NEW CENTURY", "New Century City Marketing Corp.") and pt.same_supplier("new-century city marketing corp", "New Century City Marketing Corp.")
    assert not pt.same_supplier("New Century", "AMK Auto Parts and Accessories Trading")
    parsed = pt.parse_workbook(_file(GOOD, supplier="New Century"))
    assert parsed["suppliers"] == ["New Century"] and parsed["brands"] == {"Westlake": 3, "Motolite": 1, "Enkei": 1}
    assert pt.supplier_check(nc, parsed)["level"] == "ok"
    assert pt.supplier_check(amk, parsed)["level"] == "mismatch"
    ok, msg, _ = pt.import_workbook(amk, parsed)
    assert not ok and "you chose AMK Auto Parts and Accessories Trading" in msg and "Nothing was saved" in msg
    assert fetchone(db, "SELECT COUNT(*) FROM supplier_price_list")[0] == 0
    assert pt.import_workbook(nc, parsed)[0]
    assert pt.supplier_check(nc, pt.parse_workbook(_file(GOOD)))["level"] == "none"          # an older file: a warning only


def test_a_supplier_template_names_the_supplier_even_with_no_prices_yet(db):
    nc = make_creditor(db, name="New Century City Marketing Corp.")
    wb = load_workbook(io.BytesIO(pt.build_template(nc)))
    assert wb[pt.SHEET].cell(row=2, column=1).value == "New Century City Marketing Corp."
    assert pt.parse_workbook(pt.build_template(nc))["problems"] == ["The file has no product rows."]          # the name alone is not a product row
